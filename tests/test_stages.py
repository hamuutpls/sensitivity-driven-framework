import copy
import json

import pytest
import torch

from sdf.run import start_run
from sdf.search_space import SEARCH_SPACE
from sdf.stage0.run import run_stage0
from sdf.stages.methods import METHODS, methods_for
from sdf.stages.runner import Stage0Plans, run_stage, run_stages
from conftest import fake_texts


@pytest.fixture
def stage0_dir(tiny_llama, tokenizer, small_cfg):
    ctx = start_run(small_cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16})
    run_stage0(ctx, cand, model=copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)
    return ctx.run_dir / "stage_0"


def test_registry_lists_every_spec_method():
    spec = {1: {"rtn", "gptq", "awq", "omniquant", "squeezellm", "spqr", "efficientqat", "aqlm", "quip", "quipsharp",
                "pbllm", "billm", "bitsandbytes", "qtip", "abqllm", "rtn_act", "smoothquant", "quarot", "rptq",
                "spinquant"},
            2: {"structured_prune", "unstructured_prune", "low_rank"},
            3: {"quarot_kv", "kvquant", "h2o", "snapkv", "infinigen"}}
    for stage, names in spec.items():
        assert names <= {m.name for m in METHODS.values() if m.stage == stage}
    with pytest.raises(ValueError):
        methods_for(1, ["h2o"])  # a Stage 3 method
    with pytest.raises(ValueError):
        methods_for(1, ["low_rank"])  # pruning is Stage 2


def test_series_methods_pair_a_pruning_method_with_a_stage1_weight_method():
    (m,) = methods_for(2, ["low_rank_after_gptq"])
    assert m.stage == 2 and m.quant_plans == ("quant", "quant") and m.plans == ("prune", "prune_same_size")
    assert m.calibrated and m.version != METHODS["low_rank"].version
    for bad in ("low_rank_after_rtn_act", "low_rank_after_low_rank", "rtn_after_gptq"):
        with pytest.raises(ValueError):
            methods_for(2, [bad])


def test_stage0_plans_load(stage0_dir):
    plans = Stage0Plans.load(stage0_dir)
    assert {"quant", "prune", "prune_same_size", "activations", "kv", "kv_bits_only", "joint_weights",
            "joint_acts"} <= set(plans.plans)
    assert all(lp.pruning_ratio == 0 for lp in plans.plans["quant"].layers)
    assert all(lp.bit_width == 16 for k in ("prune", "prune_same_size") for lp in plans.plans[k].layers)
    assert plans.num_layers == 4 and plans.kv_profile is not None
    with pytest.raises(FileNotFoundError):
        Stage0Plans.load(stage0_dir.parent)


ACTIVATION_METHODS_1 = ("rtn_act", "smoothquant", "quarot", "spinquant", "rptq")
FORK_WEIGHT_METHODS = ("omniquant", "squeezellm", "spqr", "efficientqat", "aqlm", "quip", "quipsharp", "pbllm", "billm",
                       "bitsandbytes")


@pytest.mark.parametrize("stage,method", [(1, "rtn"), (1, "gptq"), (1, "awq")]
                         + [(1, m) for m in ACTIVATION_METHODS_1 + FORK_WEIGHT_METHODS] + [
                                          (2, "unstructured_prune"), (2, "structured_prune"), (2, "low_rank"),
                                          (2, "unstructured_prune_after_gptq"), (2, "low_rank_after_awq"),
                                          (3, "rtn_kv")])
def test_run_stage_end_to_end(stage0_dir, tiny_llama, tokenizer, small_cfg, stage, method):
    ctx = start_run(small_cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16})
    pristine = {k: v.clone() for k, v in tiny_llama.state_dict().items()}
    spec = methods_for(stage, [method])[0]
    unimplemented = next((m.name for m in METHODS.values() if m.stage == stage and m.apply is None
                          and m.plans == methods_for(stage, [method])[0].plans), None)
    kwargs = dict(model_factory=lambda: copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)
    out = run_stage(ctx, stage, [method] + ([unimplemented] if unimplemented else []), cand, stage0_dir, **kwargs)

    res = json.loads(out["json"].read_text())
    rows = {f"{r['method']}/{r['variant']}": r for r in res["rows"]}
    assert rows["baseline/fp16"]["info"]["cached"]  # same cache entry Stage 0 measured
    assert rows[f"{method}/original"]["status"] == "ok"
    if spec.fixed_bits:  # BiLLM: no per-layer bit widths
        assert rows[f"{method}/framework"]["status"] == "failed"
        assert "cannot follow a per-layer bit plan" in rows[f"{method}/framework"]["error"]
    else:
        assert rows[f"{method}/framework"]["status"] == "ok"
    for r in rows.values():
        if r["status"] == "ok":
            assert r["metrics"]["ppl_val"] > 1 and r["metrics"]["ppl_heldout"] > 1
    # a method that is not written yet is a failed row saying so, not a crash or a silent gap
    if unimplemented:
        assert rows[f"{unimplemented}/original"]["status"] == "failed"
        assert METHODS[unimplemented].unavailable in rows[f"{unimplemented}/original"]["error"]
    # framework rows are compared with the original row of their method
    if not spec.fixed_bits:
        assert "vs_original_abs" in rows[f"{method}/framework"]["deltas"]["ppl_val"]
    assert out["report"].exists() and out["xlsx"].exists() and out["report"].parent.name == f"stage_{stage}"
    # every variant starts from a fresh model: the source model is untouched
    assert all((v == tiny_llama.state_dict()[k]).all() for k, v in pristine.items())

    if stage == 2:
        assert rows[f"{method}_same_size/framework"]["info"]["compare_to"] == method
    if stage in (1, 2) and method not in ACTIVATION_METHODS_1 and not spec.fixed_bits:
        assert "model_size_gb" not in rows[f"{method}/framework"]["metrics"]  # simulated: FP16 in memory
    if stage == 1 and method not in ACTIVATION_METHODS_1:
        assert rows[f"{method}/original"]["metrics"]["avg_bits_per_weight"] < 16
        # Stage 1 removes nothing, in any row; the zero share is reported next to the FP16 row's
        assert all(r["metrics"].get("sparsity", 0) == 0 for r in rows.values() if r["status"] == "ok")
        assert "zero_weight_share" in rows["baseline/fp16"]["metrics"]
        assert rows[f"{method}/original"]["info"]["note"] if spec.note else True
    if stage == 1 and method in ACTIVATION_METHODS_1:
        assert rows[f"{method}/original"]["metrics"]["avg_activation_bits"] < 16
    if stage == 2:
        # the standard version removes prune_ratio_aggressive from every layer; the Stage 2 framework plans keep
        # the never-pruned layers whole
        assert rows[f"{method}/original"]["metrics"]["sparsity"] > 0
        assert rows[f"{method}/framework"]["metrics"]["sparsity"] > 0
        # alone, Stage 2 leaves the weights at 16 bits; after Stage 1 they are stored at the Stage 1 bits
        bits = "bits [4]" if "_after_" in method else "bits [16]"
        assert bits in rows[f"{method}/original"]["info"]["description"]
    if stage == 3:
        assert rows["rtn_kv/framework"]["metrics"]["predicted_kv_memory_gb"] < \
            rows["baseline/fp16"]["metrics"]["predicted_kv_memory_gb"]

    # the original row is cached by config: a second run reuses it
    out2 = run_stage(ctx, stage, [method], cand, stage0_dir, **kwargs)
    rows2 = {f"{r['method']}/{r['variant']}": r for r in json.loads(out2["json"].read_text())["rows"]}
    assert rows2[f"{method}/original"]["info"]["cached"]


def test_run_stages_uses_config_methods(stage0_dir, tiny_llama, tokenizer, small_cfg):
    cfg = small_cfg.with_overrides({"stages.stage1_methods": ["rtn"], "stages.stage2_methods": ["unstructured_prune"],
                                    "stages.stage2_after": ["gptq"], "stages.stage3_methods": []})
    ctx = start_run(cfg)
    out = run_stages(ctx, SEARCH_SPACE.make({"calib_samples": 16}), stage0_dir,
                     model_factory=lambda: copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)
    assert set(out) == {f"stage_{n}_{k}" for n in (1, 2) for k in ("report", "xlsx", "json")}
    rows = json.loads((ctx.run_dir / "stage_2" / "results.json").read_text())["rows"]
    assert {r["method"] for r in rows} == {"baseline", "unstructured_prune", "unstructured_prune_same_size",
                                           "unstructured_prune_after_gptq", "unstructured_prune_after_gptq_same_size"}


def test_downstream_tasks_and_master_report(stage0_dir, tiny_llama, tokenizer, small_cfg, monkeypatch):
    from openpyxl import load_workbook

    from sdf.eval import downstream
    from sdf.reporting.master import write_master
    from sdf.reporting.metrics import METRICS

    calls = []

    def fake_accuracy(model, tok, tasks, limit, batch_size, seed):
        calls.append(tasks)
        return {**{downstream.task_metric(t): 0.5 for t in tasks}, "downstream_acc_mean": 0.5}

    monkeypatch.setattr(downstream, "downstream_accuracy", fake_accuracy)
    cfg = small_cfg.with_overrides({"eval.downstream_tasks": ["piqa"], "stages.stage1_methods": ["rtn"],
                                    "stages.stage2_methods": [], "stages.stage3_methods": []})
    ctx = start_run(cfg)
    run_stages(ctx, SEARCH_SPACE.make({"calib_samples": 16}), stage0_dir,
               model_factory=lambda: copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)
    rows = json.loads((ctx.run_dir / "stage_1" / "results.json").read_text())["rows"]
    assert all(r["metrics"]["acc_piqa"] == 0.5 for r in rows if r["status"] == "ok")
    assert METRICS["acc_piqa"].better == "higher"

    out = write_master(ctx.run_dir)  # the fixture's Stage 0 ran in the same run folder
    text = out["master_report"].read_text()
    assert "Stage 1: Quantization (weights and activations)" in text and "(lower is better)" in text and "(higher is better)" in text
    ws = load_workbook(out["all_stages_xlsx"])["All stages"]
    s0_rows = json.loads((stage0_dir / "results.json").read_text())["rows"]
    assert ws.max_row == 1 + len(s0_rows) + len(rows)
    assert "Stage 0:" in text


def test_fp16_cache_miss_builds_the_model_with_the_factory(stage0_dir, tiny_llama, tokenizer, small_cfg, tmp_path):
    ctx = start_run(small_cfg.with_overrides({"run.cache_dir": str(tmp_path / "fresh_cache")}))
    out = run_stage(ctx, 1, ["rtn"], SEARCH_SPACE.make({"calib_samples": 16}), stage0_dir,
                    model_factory=lambda: copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)
    fp16 = next(r for r in json.loads(out["json"].read_text())["rows"] if r["variant"] == "fp16")
    assert fp16["status"] == "ok" and not fp16["info"]["cached"]  # measured on the factory's model


def test_gptq_without_input_correlation_is_round_to_nearest():
    from sdf.stage0.sensitivity import round_to_nearest
    from sdf.stages.weights import gptq_quantize_

    torch.manual_seed(0)
    w = torch.randn(8, 64)
    for gs in (16, -1):
        q = w.clone()
        gptq_quantize_(q, torch.eye(64), 4, gs)  # diagonal H: no error feedback between columns
        assert torch.allclose(q, round_to_nearest(w, 4, gs, int_zero=True), atol=1e-6)


def test_gptq_beats_round_to_nearest_on_correlated_inputs():
    from sdf.stage0.sensitivity import round_to_nearest
    from sdf.stages.weights import gptq_quantize_

    torch.manual_seed(0)
    x = torch.randn(512, 8) @ torch.randn(8, 64) + 0.1 * torch.randn(512, 64)  # low-rank, correlated inputs
    w = torch.randn(16, 64)
    H = 2 * x.T @ x / len(x)
    err = lambda q: ((x @ (w - q).T) ** 2).mean()
    q = w.clone()
    gptq_quantize_(q, H, 3, 16)
    assert err(q) < 0.5 * err(round_to_nearest(w, 3, 16, int_zero=True))


def test_stage1_methods_refuse_a_plan_that_removes_anything(tiny_llama):
    from sdf.stage0.planner import uniform_plan
    from sdf.stages.weights import awq_, gptq_

    batches = [torch.randint(0, 64, (2, 16))]
    for fn in (gptq_, awq_):
        with pytest.raises(ValueError, match="removes nothing"):
            fn(copy.deepcopy(tiny_llama), uniform_plan([0.0] * 4, 4, 0.3), batches, -1, 16)


def test_awq_scales_toward_large_activations():
    from sdf.stage0.planner import uniform_plan
    from sdf.stage0.sensitivity import round_to_nearest
    from sdf.stages.weights import awq_
    from torch import nn

    class Toy(nn.Module):  # one "decoder layer": two Linears fed the same input, like q/k/v
        def __init__(self):
            super().__init__()
            self.model = nn.Module()
            self.model.layers = nn.ModuleList([nn.ModuleDict({"a": nn.Linear(64, 16, bias=False),
                                                              "b": nn.Linear(64, 8, bias=False)})])

        def forward(self, input_ids, use_cache=False):
            x = input_ids
            layer = self.model.layers[0]
            out = layer["a"](x), layer["b"](x)
            for h in layer._forward_hooks.values():
                h(layer, (x,), out)
            return out

    torch.manual_seed(0)
    x = torch.randn(256, 64)
    x[:, :4] *= 30  # a few input channels with large activations, as in real LLMs
    toy = Toy()
    w = {k: m.weight.data.clone() for k, m in toy.model.layers[0].items()}
    awq_(toy, uniform_plan([0.0], 3, 0.0), [x], 16, 16)
    for k, m in toy.model.layers[0].items():
        err = lambda q: ((x @ (w[k] - q).T) ** 2).mean()
        assert err(m.weight.data) < err(round_to_nearest(w[k], 3, 16, int_zero=True))


def test_pruning_and_low_rank_methods_remove_what_the_plan_says(tiny_llama):
    from torch import nn

    from sdf.stage0.planner import uniform_plan
    from sdf.stages.pruning import low_rank_, mlp_linears, structured_prune_, wanda_

    torch.manual_seed(0)
    batches = [torch.randint(0, 64, (2, 16)) for _ in range(2)]
    plan = uniform_plan([0.0] * 4, 16, 0.5)  # 16 bits: no rounding, only removal
    linears = lambda m: [x for x in m.modules() if isinstance(x, nn.Linear) and x is not m.lm_head]

    m = copy.deepcopy(tiny_llama)
    wanda_(m, plan, batches, -1, 16)
    assert all(abs((l.weight == 0).float().mean().item() - 0.5) < 0.05 for l in linears(m))

    m = copy.deepcopy(tiny_llama)
    structured_prune_(m, plan, batches, -1, 16)
    for layer in m.model.layers:
        ups, down = mlp_linears(layer, 32)
        dead = (down.weight == 0).all(dim=0)
        assert dead.any() and all((u.weight[dead] == 0).all() for u in ups)

    m = copy.deepcopy(tiny_llama)
    low_rank_(m, plan, batches, -1, 16)
    for l in linears(m):
        out, inp = l.weight.shape
        assert torch.linalg.matrix_rank(l.weight.float()) <= int(0.5 * out * inp / (out + inp))


def test_pruning_methods_round_nothing_except_the_low_rank_factors(tiny_llama):
    from sdf.stage0.planner import uniform_plan
    from sdf.stages.pruning import low_rank_, wanda_

    batches = [torch.randint(0, 64, (2, 16))]
    plan = uniform_plan([0.0] * 4, 4, 0.5)  # 4-bit plan: Wanda must still leave the kept weights as they were
    m = copy.deepcopy(tiny_llama)
    wanda_(m, plan, batches, -1, 16)
    for a, b in zip(m.model.layers, tiny_llama.model.layers):
        w, w0 = a.self_attn.q_proj.weight, b.self_attn.q_proj.weight
        assert torch.equal(w[w != 0], w0[w != 0])
    zero_ratio = uniform_plan([0.0] * 4, 4, 0.0)  # nothing to prune: nothing changes, even at 4 bits
    m = copy.deepcopy(tiny_llama)
    low_rank_(m, zero_ratio, batches, -1, 16)
    assert all(torch.equal(v, tiny_llama.state_dict()[k]) for k, v in m.state_dict().items())


def test_stage_settings_join_the_cache_key_only_when_changed():
    from sdf.config import Stage0Config
    from sdf.stages.runner import _setting_key

    s0 = Stage0Config()
    kinds = ("weights", "activations", "prune", "kv")
    assert all(_setting_key(k, s0) == {} for k in kinds)  # defaults: existing cache entries stay valid
    s0.weight_zero_point, s0.act_group_size = "float", 64
    assert _setting_key("weights", s0) == {"weight_zero_point": "float"} == _setting_key("prune", s0)
    assert _setting_key("activations", s0) == {"act_group_size": 64} and _setting_key("kv", s0) == {}


def test_removed_channels_and_factors_need_no_mask_in_the_predicted_size(stage0_dir, small_cfg):
    from sdf.stages.runner import plan_metrics

    plans = Stage0Plans.load(stage0_dir)
    plan = plans.plans["prune"]
    cand = SEARCH_SPACE.make({})
    size = lambda storage: plan_metrics(plan, plans, small_cfg, cand, storage)["predicted_weight_memory_gb"]
    assert size("free") < size(None) == size("bitmask")
    assert METHODS["structured_prune"].storage == METHODS["low_rank"].storage == "free"
    assert METHODS["unstructured_prune"].storage is None


def test_wanda_scores_do_not_overflow_in_fp16(tiny_llama, monkeypatch):
    from torch import nn

    from sdf.stage0.planner import uniform_plan
    from sdf.stage0.prune_sweep import magnitude_mask
    from sdf.stages import pruning

    batches = [torch.randint(0, 64, (2, 16))]
    # every channel's summed squares are 1e10 (sqrt 1e5, above the FP16 maximum of 65504): Wanda then ranks by |w|
    monkeypatch.setattr(pruning, "input_sq_norms",
                        lambda model, layer, b: {m: torch.full((m.in_features,), 1e10)
                                                 for m in layer.modules() if isinstance(m, nn.Linear)})
    m = copy.deepcopy(tiny_llama).half()
    before = [x.weight.clone() for x in m.model.layers[0].modules() if isinstance(x, nn.Linear)]
    pruning.wanda_(m, uniform_plan([0.0] * 4, 16, 0.75), batches, -1, 16)
    after = [x.weight for x in m.model.layers[0].modules() if isinstance(x, nn.Linear)]
    for b, a in zip(before, after):
        assert torch.equal(a != 0, magnitude_mask(b, 0.75))


def test_joint_methods_pair_a_weight_method_with_an_activation_method():
    m, w8 = methods_for(1, ["gptq_with_quarot", "rtn_with_rtn_act_w8a8"])
    assert m.joint == (None, None) and m.plans == ("quant", "joint_weights") and m.calibrated
    assert w8.joint == (8, 8) and w8.plans == ()
    assert methods_for(1, ["qtip_with_rtn_act"])[0].apply is None  # unported weight method: failed row
    for bad in ("rtn_act_with_rtn_act", "gptq_with_gptq", "gptq_with_h2o"):
        with pytest.raises(ValueError):
            methods_for(1, [bad])


def test_joint_rows_quantize_weights_and_activations_per_layer(stage0_dir, tiny_llama, tokenizer, small_cfg):
    from sdf.stages.methods import JointPlan, MethodCall
    from sdf.stages.runner import framework_plan, original_plan

    ctx = start_run(small_cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16})
    names = ["gptq_with_rtn_act", "gptq_with_rtn_act_w8a8", "rtn_with_quarot"]
    out = run_stage(ctx, 1, names, cand, stage0_dir, model_factory=lambda: copy.deepcopy(tiny_llama),
                    tokenizer=tokenizer, text_loader=fake_texts)
    rows = {f"{r['method']}/{r['variant']}": r for r in json.loads(out["json"].read_text())["rows"]}
    for k in ("gptq_with_rtn_act/original", "gptq_with_rtn_act/framework", "gptq_with_rtn_act_joint/framework",
              "gptq_with_rtn_act_w8a8/original", "rtn_with_quarot/original", "rtn_with_quarot/framework",
              "rtn_with_quarot_joint/framework"):
        assert rows[k]["status"] == "ok", rows[k].get("error")
        assert rows[k]["metrics"]["avg_activation_bits"] < 16 and rows[k]["metrics"]["avg_bits_per_weight"] < 16
    assert "gptq_with_rtn_act_w8a8/framework" not in rows
    assert "'W4A8': 4" in rows["gptq_with_rtn_act/original"]["info"]["description"]
    assert "'W8A8': 4" in rows["gptq_with_rtn_act_w8a8/original"]["info"]["description"]
    assert rows["gptq_with_rtn_act/framework"]["deltas"]["ppl_val"]["vs_original_abs"] is not None

    # the framework plan is the Stage 0 weight plan and activation plan, layer by layer
    plans = Stage0Plans.load(stage0_dir)
    m = methods_for(1, ["gptq_with_rtn_act"])[0]
    plan = framework_plan(m, plans, "quant")
    assert isinstance(plan, JointPlan) and plan.weights == plans.plans["quant"] and plan.acts == plans.plans["activations"]
    assert original_plan(m, plans, small_cfg.stage0, 0.0).combos() == {"W4A8": 4}
    jp = framework_plan(m, plans, "joint_weights")
    assert jp.weights == plans.plans["joint_weights"] and jp.acts == plans.plans["joint_acts"]
    assert rows["gptq_with_rtn_act_joint/framework"]["info"]["compare_to"] == "gptq_with_rtn_act"
    # while measured, every Linear of a layer with activation bits < 16 rounds its input
    model = copy.deepcopy(tiny_llama)
    call = MethodCall(model, plan, cand, small_cfg, lambda: [torch.randint(0, tiny_llama.config.vocab_size, (2, 16))])
    from sdf.stage0.sensitivity import find_decoder_layers
    with m.apply(call):
        hooked = [any(l._forward_pre_hooks for l in layer.modules() if isinstance(l, torch.nn.Linear))
                  for layer in find_decoder_layers(model)]
    assert hooked == [lp.act_bits < 16 for lp in plan.acts.layers]
