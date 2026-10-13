"""Fixes from the 2026-10-13 ponytail audit (numbers are the audit's findings). Every test here fails on main 9a7ef8f."""

import copy
import json
import math
import subprocess
import sys
import types
from pathlib import Path

import pytest
import torch
from openpyxl import load_workbook

from sdf.config import FrameworkConfig, Stage0Config
from sdf.data import eval_data_label
from sdf.reporting.markdown import _plain_verdicts
from sdf.reporting.reporter import StageReporter
from sdf.requirements import DeploymentRequirement
from sdf.stage0.planner import same_size_pruning_plan, uniform_plan
from sdf.stages.methods import METHODS, methods_for
from sdf.stages.pruning import mlp_linears, structured_prune_


# 1 ------------------------------------------------------------------------------------------------------------
def test_same_size_plan_at_ratio_zero_prunes_nothing():
    plan = same_size_pruning_plan([0.0, 0.3, 0.6, 1.0], 16, 0.0, [100] * 4, frozenset({3}))
    assert [lp.pruning_ratio for lp in plan.layers] == [0.0] * 4


# 2 ------------------------------------------------------------------------------------------------------------
def test_structured_pruning_refuses_a_share_it_cannot_remove(tiny_llama):
    # tiny model: feed-forward weights are 3 x 32 x 64 = 6144 of 9216 in a layer (2/3)
    batches = [torch.randint(0, 64, (2, 16))]
    with pytest.raises(ValueError, match="at most 66.7%"):
        structured_prune_(copy.deepcopy(tiny_llama), uniform_plan([0.0] * 4, 16, 0.9), batches, -1, 16)
    m = copy.deepcopy(tiny_llama)
    structured_prune_(m, uniform_plan([0.0] * 4, 16, 0.5), batches, -1, 16)
    for layer in m.model.layers:
        _, down = mlp_linears(layer, 32)
        assert int((down.weight == 0).all(dim=0).sum()) == round(0.5 * 9216 / 96)  # 48 of 64 channels


# 3, 13 ------------------------------------------------------------------------------------------------------
class _Stop(Exception):
    pass


def _fake_vllm(monkeypatch, generate):
    seen = {}

    class LLM:
        def __init__(self, **kw):
            seen.update(kw)

        def apply_model(self, fn):
            raise _Stop

        def generate(self, *a, **kw):
            return generate(*a, **kw)

    monkeypatch.setitem(sys.modules, "vllm", types.SimpleNamespace(LLM=LLM, SamplingParams=lambda **k: k))
    monkeypatch.setattr(torch.cuda, "mem_get_info", lambda: (40e9, 40e9))
    monkeypatch.setattr(torch.cuda, "synchronize", lambda: None)
    monkeypatch.setattr(torch.cuda, "max_memory_allocated", lambda: 0)
    return seen


def test_vllm_worker_turns_prefix_caching_off_and_fits_the_latency_run(monkeypatch):
    from sdf.eval import vllm_worker

    def generate(prompts, sp, use_tqdm):
        raise _Stop

    seen = _fake_vllm(monkeypatch, generate)
    with pytest.raises(_Stop):
        vllm_worker.main({"model_dir": "m", "seq_len": 128, "prompt_len": 128, "decode_tokens": 64,
                          "linear_backend": None, "val": [], "heldout": []})
    assert seen["enable_prefix_caching"] is False
    assert seen["max_model_len"] == 128 + 64 + 8  # was seq_len + 8 = 136: decoding stopped early


def test_vllm_perplexity_pairs_each_token_with_its_own_logprob(monkeypatch):
    from sdf.eval import vllm_worker

    def generate(prompts, sp, use_tqdm):
        out = []
        for p in prompts:
            # position i predicts token i with log-probability -i (position 0 has none)
            lps = [None] + [{tok: types.SimpleNamespace(logprob=-float(i))} for i, tok in
                            enumerate(p["prompt_token_ids"]) if i > 0]
            out.append(types.SimpleNamespace(prompt_logprobs=lps))
        return out

    _fake_vllm(monkeypatch, generate)
    windows = [[5, 6, 7, 8]]
    out = vllm_worker.main({"model_dir": "m", "seq_len": 4, "prompt_len": 2, "decode_tokens": 2,
                            "linear_backend": None, "val": windows, "heldout": windows, "warmup": 0, "repeats": 1})
    assert math.isclose(out["vllm_ppl_val"], math.exp((1 + 2 + 3) / 3), rel_tol=1e-9)  # mean of 1, 2, 3
    with pytest.raises(ValueError, match="longer than an eval window"):  # a prompt is never silently cut
        vllm_worker.main({"model_dir": "m", "seq_len": 4, "prompt_len": 8, "decode_tokens": 2,
                          "linear_backend": None, "val": windows, "heldout": windows, "warmup": 0, "repeats": 1})


def _fake_trt(monkeypatch, tmp_path, logits):
    from sdf.eval import trtllm_worker

    built = {}

    class LLM:
        def __init__(self, **kw):
            built.update(kw)

        def save(self, d):
            Path(d).mkdir(parents=True, exist_ok=True)

        def generate(self, windows, sp, use_tqdm):
            return [types.SimpleNamespace(context_logits=logits(len(w))) for w in windows]

    fakes = {"tensorrt_llm": types.ModuleType("tensorrt_llm"),
             "tensorrt_llm._tensorrt_engine": types.SimpleNamespace(LLM=LLM),
             "tensorrt_llm.llmapi": types.SimpleNamespace(BuildConfig=lambda **k: k, KvCacheConfig=lambda **k: k,
                                                          SamplingParams=lambda **k: k)}
    for k, v in fakes.items():
        monkeypatch.setitem(sys.modules, k, v)
    monkeypatch.setattr(trtllm_worker, "_used_gb", lambda: 0.0)
    monkeypatch.setattr(torch.cuda, "synchronize", lambda: None)
    hf = tmp_path / "hf"
    hf.mkdir(exist_ok=True)
    return trtllm_worker, built, hf


def test_trtllm_worker_turns_block_reuse_off_and_reads_logits_at_the_right_position(monkeypatch, tmp_path):
    # row j of the logits predicts token j + 1: put all the mass on token (j + 1) at row j, so a shift by one
    # position gives a huge perplexity instead of 1
    def logits(n):
        x = torch.full((n, 8), -1e4)
        for j in range(n - 1):
            x[j, j + 2] = 0.0  # window [1, 2, 3, 4]: token at j + 1 is j + 2
        return x

    worker, built, hf = _fake_trt(monkeypatch, tmp_path, logits)
    out = worker.main({"model_dir": str(hf), "engine_dir": str(tmp_path / "engine"), "seq_len": 4,
                       "val": [[1, 2, 3, 4]], "heldout": [[1, 2, 3, 4]], "prompt_len": 2, "warmup": 0, "repeats": 1,
                       "decode_tokens": 6})
    assert math.isclose(out["trtllm_ppl_val"], 1.0, rel_tol=1e-6)
    assert built["kv_cache_config"]["enable_block_reuse"] is False
    assert built["build_config"]["max_seq_len"] == max(4, 2 + 6) + 8
    with pytest.raises(ValueError, match="longer than an eval window"):  # checked before the engine is built
        worker.main({"model_dir": str(hf), "engine_dir": str(tmp_path / "e2"), "seq_len": 4, "val": [[1, 2, 3, 4]],
                     "heldout": [[1, 2, 3, 4]], "prompt_len": 8, "warmup": 0, "repeats": 1, "decode_tokens": 2})


# 4 ------------------------------------------------------------------------------------------------------------
def _reporter(tmp_path, metrics):
    return StageReporter(1, tmp_path, "t", {}, {}, {}, DeploymentRequirement(), list(metrics))


def _pair(rep, original, framework):
    with rep.method("m", "original") as r:
        r.metrics.update(original)
    with rep.method("m", "framework") as r:
        r.metrics.update(framework)
    rep.compute_deltas()


def test_findings_and_plain_summary_give_the_same_verdict_and_noise_is_a_tie(tmp_path):
    rep = _reporter(tmp_path, ["ppl_val", "decode_ms_per_token_mean", "peak_memory_gb", "build_time_s"])
    # the audit's case: better perplexity, 0.001 ms faster (std 0.2), 1% more peak memory
    _pair(rep, {"ppl_val": 10.0, "decode_ms_per_token_mean": 5.000, "decode_ms_per_token_std": 0.2,
                "peak_memory_gb": 1.00, "decode_tokens_per_s": 200.0, "build_time_s": 10.0},
          {"ppl_val": 9.0, "decode_ms_per_token_mean": 4.999, "decode_ms_per_token_std": 0.2,
           "peak_memory_gb": 1.01, "decode_tokens_per_s": 200.04, "build_time_s": 25.0})
    technical, plain = rep.comparison_summary()[0], _plain_verdicts(rep)[0]
    assert "trades off" in technical and "trade-off" in plain
    assert "Decode latency" not in technical and "throughput" not in technical.lower()  # noise; not a main metric
    assert rep.verdict(rep.rows[1]) == (["ppl_val"], ["peak_memory_gb"], ["build_time_s"])
    assert "One-off cost: Build time (+15" in technical and "paid once" in plain


# 5, 6 ---------------------------------------------------------------------------------------------------------
def test_series_methods_keep_the_stage1_limits():
    qtip = methods_for(2, ["unstructured_prune_after_qtip"])[0]
    assert qtip.apply is None and "QTIP" in qtip.unavailable
    billm = methods_for(2, ["unstructured_prune_after_billm"])[0]
    assert billm.apply is not None and billm.fixed_bits == METHODS["billm"].fixed_bits != ""
    low_rank = methods_for(2, ["low_rank_after_gptq"])[0]
    assert low_rank.apply is None and "replaces the Stage 1 rounding" in low_rank.unavailable
    assert methods_for(2, ["unstructured_prune_after_gptq"])[0].apply is not None


# 7 ------------------------------------------------------------------------------------------------------------
def test_stage0_removes_plan_files_of_an_earlier_run(tiny_llama, tokenizer, small_cfg):
    from conftest import fake_texts
    from sdf.run import start_run
    from sdf.search_space import SEARCH_SPACE
    from sdf.stage0.run import run_stage0

    cfg = small_cfg.with_overrides({"stage0.joint_plan": False})
    ctx = start_run(cfg)
    stale = ctx.run_dir / "stage_0" / "joint_weight_plan.json"
    stale.parent.mkdir(parents=True, exist_ok=True)
    stale.write_text("{}")  # left by an earlier run with the joint plan on
    run_stage0(ctx, SEARCH_SPACE.make({"calib_samples": 16}), model=tiny_llama, tokenizer=tokenizer,
               text_loader=fake_texts)
    assert not stale.exists() and (ctx.run_dir / "stage_0" / "quant_plan.json").exists()


# 9 ------------------------------------------------------------------------------------------------------------
def test_fp16_key_ignores_vllm_and_trtllm_settings(tmp_path):
    from sdf.run import start_run
    from sdf.stage0.run import fp16_key

    base = FrameworkConfig().with_overrides({"run.output_root": str(tmp_path), "run.run_id": "k"})
    other = base.with_overrides({"eval.vllm_python": "/x/python", "eval.vllm_gpu_memory_utilization": 0.5,
                                 "eval.trtllm_python": "/y/python"})
    cpu = torch.device("cpu")
    assert fp16_key(start_run(base), cpu) == fp16_key(start_run(other), cpu)
    assert fp16_key(start_run(base), cpu) != fp16_key(start_run(base.with_overrides({"eval.seq_len": 256})), cpu)


# 10 -----------------------------------------------------------------------------------------------------------
def test_kv_uniform_bits_must_be_a_measured_option():
    with pytest.raises(ValueError, match="kv_uniform_bits 3"):
        Stage0Config(kv_uniform_bits=3)
    assert Stage0Config(kv_uniform_bits=16).kv_uniform_bits == 16  # at the baseline: nothing rounded


# 14 -----------------------------------------------------------------------------------------------------------
def test_llamacpp_refuses_a_sentencepiece_tokenizer_without_its_model_file(tmp_path):
    from sdf.eval.llamacpp import _save_tokenizer

    src = tmp_path / "tok"
    src.mkdir()

    class Tok:
        name_or_path = str(src)

        def __init__(self, files):
            self.vocab_files_names = files

        def save_pretrained(self, out):
            Path(out).mkdir(parents=True, exist_ok=True)

    with pytest.raises(FileNotFoundError, match="tokenizer.model"):
        _save_tokenizer(Tok({"vocab_file": "tokenizer.model"}), tmp_path / "a")
    _save_tokenizer(Tok({"vocab_file": "vocab.json"}), tmp_path / "b")  # no SentencePiece file expected: fine
    (src / "tokenizer.model").write_bytes(b"sp")
    _save_tokenizer(Tok({"vocab_file": "tokenizer.model"}), tmp_path / "c")
    assert (tmp_path / "c" / "tokenizer.model").read_bytes() == b"sp"


# 15 -----------------------------------------------------------------------------------------------------------
def test_rtn_act_rounds_by_the_same_rule_alone_and_in_a_joint_method(tiny_llama):
    from sdf.stage0.activation import uniform_activation_plan
    from sdf.stages.methods import WITH, JointPlan, MethodCall, _rtn_activations

    x = torch.randint(0, 64, (1, 8))
    weights16 = uniform_plan([0.0] * 4, 16, 0.0)  # weights left alone: only the activation rounding differs
    joint = methods_for(1, [f"rtn{WITH}rtn_act"])[0]

    def logits(cfg, bits):
        acts = uniform_activation_plan(4, bits)
        with torch.no_grad():
            with _rtn_activations(MethodCall(tiny_llama, acts, {"gptq_groupsize": 128}, cfg, lambda: [])):
                alone = tiny_llama(x).logits
            with joint.apply(MethodCall(tiny_llama, JointPlan(weights16, acts), {"gptq_groupsize": 128}, cfg,
                                        lambda: [])):
                together = tiny_llama(x).logits
        return alone, together

    with torch.no_grad():
        ref = tiny_llama(x).logits
    base8 = FrameworkConfig().with_overrides({"stage0.baseline_bits": 8})
    alone, together = logits(base8, 8)  # at the baseline: nothing rounded, alone or joint (joint rounded before)
    assert torch.equal(alone, ref) and torch.equal(together, ref)
    alone, together = logits(base8, 4)
    assert torch.equal(alone, together) and not torch.equal(alone, ref)


# 18 -----------------------------------------------------------------------------------------------------------
def test_methods_that_ignore_the_group_size_do_not_key_their_cache_on_it():
    for name in ("squeezellm", "spqr", "aqlm", "quipsharp", "pbllm", "billm"):
        assert "gptq_groupsize" not in METHODS[name].params, name
    for name in ("gptq", "awq", "omniquant", "efficientqat", "quip"):
        assert "gptq_groupsize" in METHODS[name].params, name


# 19 -----------------------------------------------------------------------------------------------------------
def test_master_workbook_shows_nan_instead_of_a_blank_cell(tmp_path):
    from sdf.reporting.master import write_master

    (tmp_path / "stage_1").mkdir()
    row = {"method": "m", "variant": "framework", "status": "ok", "metrics": {"ppl_val": math.nan}}
    (tmp_path / "stage_1" / "results.json").write_text(json.dumps({"stage": 1, "title": "t", "rows": [row]}))
    ws = load_workbook(write_master(tmp_path)["all_stages_xlsx"])["All stages"]
    assert "nan" in {c.value for r in ws.iter_rows() for c in r}


# 20 -----------------------------------------------------------------------------------------------------------
def test_eval_data_label_follows_the_dataset_setting():
    ev = FrameworkConfig().eval
    assert eval_data_label(ev) == "wikitext-2 test, 512-token windows, validation/held-out halves"
    ev = FrameworkConfig().with_overrides({"eval.dataset": "c4", "eval.seq_len": 256}).eval
    assert eval_data_label(ev).startswith("c4 test, 256-token windows")


def test_sweep_grid_refuses_a_setting_stage0_ignores():
    from sdf.cli import sweep_main
    from sdf.stage0.sweep import default_grid

    assert "smoothquant_alpha" not in default_grid() and "gptq_groupsize" in default_grid()
    with pytest.raises(SystemExit, match="smoothquant_alpha"):
        sweep_main(["--grid", "smoothquant_alpha=0.3,0.5"])


# lean: run_worker ---------------------------------------------------------------------------------------------
def test_run_worker_returns_the_last_json_line(monkeypatch, tmp_path):
    from sdf.eval.worker import run_worker

    monkeypatch.setattr(subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(
        cmd, 0, 'log line\n{"x_ppl_val": 8.5}\n', ""))
    assert run_worker("python", "vllm_worker.py", {"a": 1}, tmp_path, RuntimeError) == {"x_ppl_val": 8.5}
    assert json.loads((tmp_path / "args.json").read_text()) == {"a": 1}
