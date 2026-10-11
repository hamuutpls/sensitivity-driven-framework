"""Edge cases found by the 2026-10-11 test sweep: each test failed before its fix."""

import logging
import math

import pytest
import torch
from openpyxl import load_workbook

from sdf.config import FrameworkConfig
from sdf.reporting.markdown import _plain_verdicts
from sdf.reporting.reporter import StageReporter
from sdf.requirements import DeploymentRequirement
from sdf.stage0 import sensitivity
from sdf.stage0.joint import JointProfile, plan_joint
from sdf.stage0.kv_cache import allocate_bits
from sdf.stage0.planner import predict_cost, uniform_plan
from sdf.stage0.sensitivity import SensitivityProfile, _ppl
from sdf.stage0.sweep import _spearman
from sdf.utils.logging import setup_logging


def _reporter(tmp_path, metrics=("ppl_val",)):
    return StageReporter(1, tmp_path, "t", {}, {}, {}, DeploymentRequirement(), list(metrics))


def _pair(rep, original, framework):
    with rep.method("m", "original") as r:
        r.metrics.update(original)
    with rep.method("m", "framework") as r:
        r.metrics.update(framework)
    rep.compute_deltas()


# --- reporting ---------------------------------------------------------------------------------------------------

def test_equal_results_read_as_a_tie_not_as_missing(tmp_path):
    rep = _reporter(tmp_path)
    _pair(rep, {"ppl_val": 10.0}, {"ppl_val": 10.0})
    assert rep.comparison_summary() == ["**m**: the framework ties with the original method."]


def test_comparison_summary_names_beats_and_loses(tmp_path):
    rep = _reporter(tmp_path)
    _pair(rep, {"ppl_val": 10.0}, {"ppl_val": 9.0})
    assert "beats" in rep.comparison_summary()[0]
    rep = _reporter(tmp_path)
    _pair(rep, {"ppl_val": 10.0}, {"ppl_val": 11.0})
    assert "loses to" in rep.comparison_summary()[0]


def test_percent_change_keeps_its_sign_for_a_negative_reference(tmp_path):
    rep = _reporter(tmp_path, ["predicted_act_ppl_rise"])
    _pair(rep, {"predicted_act_ppl_rise": -0.01}, {"predicted_act_ppl_rise": 0.01})
    d = rep.rows[1].deltas["predicted_act_ppl_rise"]
    assert d["vs_original_abs"] == pytest.approx(0.02) and d["vs_original_pct"] == pytest.approx(200.0)


def test_plain_verdict_skips_a_nan_measurement(tmp_path):
    rep = _reporter(tmp_path)
    _pair(rep, {"ppl_val": math.nan}, {"ppl_val": 12.0})
    assert not any("lost" in v for v in _plain_verdicts(rep))


def test_workbook_shows_nan_and_inf_instead_of_blank_cells(tmp_path):
    rep = _reporter(tmp_path)
    with rep.method("base", "fp16") as r:
        r.metrics.update(ppl_val=7.0)
    _pair(rep, {"ppl_val": math.nan}, {"ppl_val": math.inf})
    ws = load_workbook(rep.finalize()["xlsx"])["Summary"]
    values = {c.value for row in ws.iter_rows() for c in row}
    assert {"nan", "inf"} <= values


def test_diverged_run_fails_the_accuracy_target():
    for bad in (math.inf, math.nan):
        check = DeploymentRequirement(target_ppl=20).check({"ppl_val": bad})
        assert check.met is False and check.shortfall["target_ppl"] == math.inf
    assert DeploymentRequirement(target_ppl=20).check({}).met is None  # still "not measured" when absent


# --- config --------------------------------------------------------------------------------------------------------

def test_overrides_are_cast_to_the_field_type():
    cfg = FrameworkConfig().with_overrides({"run.run_id": 20261005, "eval.max_windows": "1e3",
                                            "stage0.act_avg_bits": 6})
    assert cfg.run.run_id == "20261005" and cfg.eval.max_windows == 1000
    assert isinstance(cfg.stage0.act_avg_bits, float)
    with pytest.raises(ValueError):
        FrameworkConfig().with_overrides({"eval.max_windows": "many"})
    with pytest.raises(ValueError):
        FrameworkConfig().with_overrides({"eval.max_windows": 1.5})


@pytest.mark.parametrize("name", ["act_bits_options", "kv_bits_options", "joint_w_options", "joint_a_options"])
def test_bit_options_must_be_ascending(name):
    with pytest.raises(ValueError, match="ascending"):
        FrameworkConfig().with_overrides({f"stage0.{name}": [8, 4]})


def test_second_run_log_does_not_go_into_the_first(tmp_path):
    a, b = tmp_path / "a" / "run.log", tmp_path / "b" / "run.log"
    try:
        setup_logging(log_file=a)
        setup_logging(log_file=b)
        logging.getLogger("sdf.test").info("second run line")
        for h in logging.getLogger("sdf").handlers:
            h.flush()
        assert "second run line" in b.read_text() and "second run line" not in a.read_text()
    finally:
        root = logging.getLogger("sdf")
        for h in [h for h in root.handlers if isinstance(h, logging.FileHandler)]:
            root.removeHandler(h)
            h.close()


def test_cache_keys_change_with_the_data_source_only_when_it_differs(small_cfg):
    from types import SimpleNamespace

    from sdf.stage0.run import profile_key, source_key

    cand = {"calib_dataset": "wikitext2", "calib_samples": 8}
    assert source_key(small_cfg, "wikitext2") == {}  # default source: old cache entries keep their keys
    other = small_cfg.with_overrides({"data.sources": {**small_cfg.data.sources,
                                                       "wikitext2": {"path": "x", "splits": {"train": "train"}}}})
    assert profile_key(SimpleNamespace(cfg=small_cfg), cand) != profile_key(SimpleNamespace(cfg=other), cand)


# --- Stage 0 -------------------------------------------------------------------------------------------------------

def test_joint_plan_keeps_the_exact_weight_budget_with_uneven_layers():
    numel = [10, 14]
    profile = JointProfile([4, 8], [16], [[[0.1], [0.0]], [[5.0], [0.0]]], numel)
    picks = plan_joint(profile, avg_weight_bits=6.0, avg_act_bits=16.0)
    assert sum(w * n for (w, _), n in zip(picks, numel)) / sum(numel) <= 6.0


def test_bit_allocation_refuses_a_budget_below_the_cheapest_option():
    with pytest.raises(ValueError, match="below the smallest option"):
        allocate_bits({0: (1, [1.0, 0.0])}, [4, 8], 3.0)
    assert allocate_bits({0: (1, [1.0, 0.0]), 1: (1, [0.5, 0.0])}, [4, 8], 6.0) == {0: 8, 1: 4}


def test_rank_agreement_with_a_flat_profile_is_undefined():
    assert math.isnan(_spearman([1.0, 1.0, 1.0], [1.0, 2.0, 3.0]))
    assert _spearman([1.0, 2.0, 3.0], [2.0, 4.0, 9.0]) == pytest.approx(1.0)


def test_huge_loss_gives_infinite_perplexity_not_a_crash(tiny_llama, monkeypatch):
    assert _ppl(800.0) == math.inf and _ppl(0.0) == 1.0
    calls = iter([1.0] + [800.0] * 4)  # baseline, then every removed layer explodes
    monkeypatch.setattr(sensitivity, "_mean_loss", lambda *a: next(calls))

    prof = sensitivity.profile_by_ablation(tiny_llama, [torch.zeros(1, 4, dtype=torch.long)], "layer_removal", "cpu")
    assert all(math.isfinite(s) and s > 0 for s in prof.raw_scores)


def test_predicted_size_refuses_mismatched_layer_counts():
    plan = uniform_plan([0.0, 0.5, 1.0], 4, 0.0)
    no_rows = SensitivityProfile([1.0, 2.0, 3.0], layer_numel=[100] * 3, layer_rows=[], other_numel=50)
    with pytest.raises(ValueError, match="layer counts differ"):
        predict_cost(plan, no_rows, 100, 32, 16)
    full = SensitivityProfile([1.0, 2.0, 3.0], layer_numel=[100] * 3, layer_rows=[10] * 3, other_numel=50)
    with pytest.raises(ValueError, match="layer counts differ"):
        predict_cost(plan, full, 100, 32, 16, reference_scores=[0.5])


def test_prune_levels_that_share_a_row_name_are_rejected():
    from sdf.stage0.prune_sweep import _validated

    with pytest.raises(ValueError, match="at least 1%"):
        _validated([0.1, 0.104])
    assert _validated([0.2, 0.1]) == [0.1, 0.2]


# --- serving export and eval ---------------------------------------------------------------------------------------

class _Tok:
    def save_pretrained(self, d):
        pass


class _Block(torch.nn.Module):
    def __init__(self):
        super().__init__()
        self.proj = torch.nn.Linear(128, 64, bias=False)


class _Model(torch.nn.Module):
    config = type("C", (), {"to_dict": lambda self: {"model_type": "llama"}})()

    def __init__(self):
        super().__init__()
        self.model = torch.nn.Module()
        self.model.layers = torch.nn.ModuleList([_Block() for _ in range(3)])
        self.lm_head = torch.nn.Linear(128, 64, bias=False)


def _plan_with(bits_by_layer):
    from dataclasses import replace

    plan = uniform_plan([0.0] * 3, 4, 0.0)
    return replace(plan, layers=tuple(replace(lp, bit_width=bits_by_layer.get(lp.layer, 4)) for lp in plan.layers))


def test_per_channel_group_size_exports_one_group_per_row():
    from sdf.eval.gptq_export import quantize_linear
    from sdf.search_space import PER_CHANNEL

    t = quantize_linear(torch.randn(64, 256), 4, PER_CHANNEL)
    assert t["scales"].shape == (1, 64) and int(t["g_idx"].max()) == 0


def test_a_16_bit_layer_is_marked_unquantized_in_both_serving_configs(tmp_path):
    import json

    from sdf.eval.gptq_export import export
    from sdf.eval.trtllm_backend import quant_file

    plan = _plan_with({1: 16})
    export(_Model(), _Tok(), plan, 128, tmp_path)
    q = json.loads((tmp_path / "config.json").read_text())["quantization_config"]
    assert q["bits"] == 4 and q["dynamic"] == {r"-:.*layers\.1\..*": {}}
    layers = quant_file(_Model(), plan, 128)["quantization"]["quantized_layers"]
    assert "model.layers.1.proj" not in layers and "W16A16_GPTQ" not in json.dumps(layers)


def test_a_plan_the_export_cannot_hold_is_a_backend_error(monkeypatch):
    from sdf.config import EvalConfig
    from sdf.eval import trtllm_backend, vllm_backend

    ids = torch.zeros(1, 4, dtype=torch.long)
    for mod, err in ((vllm_backend, vllm_backend.VLLMError), (trtllm_backend, trtllm_backend.TRTLLMError)):
        with pytest.raises(err, match="GPTQ export failed"):  # 6 bits: no GPTQ packing exists
            mod.measure(_Model(), _Tok(), uniform_plan([0.0] * 3, 6, 0.0), ids, ids, EvalConfig(), 16, 128)


def test_perplexity_of_a_diverged_model_raises_the_documented_error(monkeypatch):
    from types import SimpleNamespace

    from sdf.eval.metrics import perplexity

    class Exploded(torch.nn.Module):
        def forward(self, input_ids, labels):
            return SimpleNamespace(loss=torch.tensor(800.0))

    with pytest.raises(FloatingPointError):
        perplexity(Exploded(), torch.zeros(2, 4, dtype=torch.long), torch.device("cpu"))


def test_integer_zero_grid_reaches_the_max_of_an_all_positive_group():
    from sdf.stage0.sensitivity import round_to_nearest

    w = torch.tensor([[1.0, 2.0, 3.0, 4.0] * 4])
    q = round_to_nearest(w, 4, 16, True)
    assert q.max() == pytest.approx(4.0) and (q - w).abs().max() <= 4.0 / 15 / 2 + 1e-6
