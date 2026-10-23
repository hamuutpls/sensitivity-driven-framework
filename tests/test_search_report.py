"""The settings search's report, written from synthetic tries (no model is built)."""

import json
import random
import re
from pathlib import Path

import openpyxl
import pytest

from sdf.cli import search_report_main
from sdf.reporting.metrics import label
from sdf.reporting.search import (_tokens_per_s, best_panels, best_so_far, front_labels, goal_callouts,
                                  measured_text)
from sdf.search import FAILED, _sample, pareto, technique_space

OBJ = ["ppl_val", "vllm_decode_ms_per_token"]


@pytest.fixture
def run_dir(small_cfg):
    cfg = small_cfg.with_overrides({"requirement.target_memory_gb": 1.1, "search.objectives": OBJ,
                                    "search.stage1_methods": ["rtn_with_rtn_act", "gptq_with_smoothquant"],
                                    "search.stage3_methods": ["rtn_kv", "h2o"]})
    run = Path(small_cfg.run.output_root) / "20261021-colab-a100-search"
    (run / "search").mkdir(parents=True)
    (run / "config.json").write_text(json.dumps(cfg.to_dict()))
    space, rng, trials = technique_space(cfg), random.Random(0), []
    for name in cfg.search.searchers:
        for n in range(1, 13):
            size = rng.uniform(0.7, 1.4)
            over, failed = size > 1.1, n == 5
            rough = name == "mfbo" and n % 3
            ok = not over and not failed
            metrics = {} if over or failed else {
                "ppl_val": rng.uniform(8, 14), "ppl_heldout": rng.uniform(8, 14), "predicted_weight_memory_gb": size,
                "vllm_decode_ms_per_token": rng.uniform(4, 9), "vllm_ppl_val": rng.uniform(8, 14),
                "avg_bits_per_weight": 4.5, "avg_activation_bits": 6.0}
            trials.append({
                "searcher": name, "n": n, "point": _sample(space, rng), "windows": 2 if rough else 6,
                "cost": 0.0 if over else (1 / 3 if rough else 1.0), "built": not over, "over_limit": over,
                "full": not rough, "ok": ok, "met": ok, "values": [metrics[o] for o in OBJ] if ok else [FAILED] * 2,
                "violation": 0.0 if ok else 1.0, "metrics": metrics, "dir": "x", "time_s": 60.0, "cached": False,
                "error": "plan over the limit" if over else ("CUDA out of memory" if failed else None)})
    (run / "search" / "trials.json").write_text(json.dumps(trials))
    (run / "search" / "search_report.md").write_text("the report the search wrote when it ended")  # a finished run
    return run, trials


def test_the_report_is_rebuilt_from_saved_tries(run_dir, capsys):
    run, trials = run_dir
    search_report_main([str(run), "--tests", "417 passed, 3 skipped at 5f440dc"])
    assert "search_report:" in capsys.readouterr().out
    d = run / "search"
    wb = openpyxl.load_workbook(d / "search_results.xlsx")
    assert wb.sheetnames == ["Trials", "Pareto", "Searcher comparison", "Best configs"]
    assert wb["Trials"].max_row == len(trials) + 1
    assert wb["Pareto"].max_row == len(pareto(trials)) + 1
    text = (d / "search_report.md").read_text()
    assert "417 passed, 3 skipped at 5f440dc" in text
    heads = re.findall(r"^## (.+)$", text, re.M)
    assert heads[:4] == ["Summary", "1. Best trade-off settings", "2. Search methods compared",
                         "3. Best setting for each goal"]
    assert heads[-2:] == ["Appendix A: the other settings of each best trade-off option",
                          "Appendix B: search methods, every column"]
    assert text.count("> **What this shows.**") == text.count("> **What it means.**") == 4
    for png in re.findall(r"!\[[^]]*\]\(([^)]+)\)", text):  # every chart the report shows exists
        assert (d / png).exists(), png
    assert {"pareto_ppl_vs_predicted_weight_memory_gb.png", "pareto_ppl_vs_vllm_decode_ms_per_token.png",
            "best_so_far.png", "hypervolume.png", "best_settings.png"} <= {p.name for p in d.glob("*.png")}
    assert "](best_settings.png)" in text
    # the most accurate trade-off setting's speed is also given in tokens per second, next to its time
    top = min(pareto(trials), key=lambda t: t["values"][0])
    ms = top["metrics"]["vllm_decode_ms_per_token"]
    row = next(r for r in text.splitlines() if r.startswith("| 1 |"))
    assert f"{ms:.4g} | {1000 / ms:.4g}" in row and "| W4.5A6 |" in row
    assert "Same speed, tokens per second, higher is better" in text
    assert "predicted weight memory at most 1.1 GB" in text
    assert "2. Search methods compared\n\n> **What this shows.** How well each search method used the same budget: " \
           "the best perplexity (validation half) it found" in text
    # axis notes follow the real limits: perplexity, size and the first try's budget are above zero, hypervolume is not
    assert text.count("Neither axis starts at zero.") == 3
    hv = next(r for r in text.splitlines() if r.startswith("*How much of the space"))
    assert "zero" not in hv


def test_a_rebuild_never_rewrites_the_saved_tries_and_keeps_out_of_an_unfinished_run(run_dir):
    run, _ = run_dir
    d = run / "search"
    (d / "search_report.md").unlink()  # the search is still running
    saved = (d / "trials.json").read_bytes()
    search_report_main([str(run)])
    assert sorted(p.name for p in d.iterdir()) == ["partial", "trials.json"]
    assert (d / "partial" / "search_report.md").exists() and not (d / "partial" / "trials.json").exists()
    (d / "search_report.md").write_text("the report the search wrote when it ended")
    search_report_main([str(run)])  # finished: rebuilt in place
    assert (d / "search_report.md").read_text().startswith("# Settings search")
    assert (d / "trials.json").read_bytes() == saved


def test_a_resumed_search_whose_report_is_older_than_its_tries_counts_as_unfinished(run_dir):
    import os

    run, _ = run_dir
    d = run / "search"
    old = (d / "search_report.md").stat().st_mtime - 60
    os.utime(d / "search_report.md", (old, old))  # the last run's report; the resumed search rewrote trials.json
    search_report_main([str(run)])
    assert (d / "search_report.md").read_text() == "the report the search wrote when it ended"
    assert (d / "partial" / "search_report.md").exists()


def test_more_search_methods_than_colours_fail_loudly(run_dir):
    run, trials = run_dir
    extra = [{**trials[0], "searcher": f"s{i}"} for i in range(5)]  # 3 + 5 = 8 methods, 7 colours
    (run / "search" / "trials.json").write_text(json.dumps(trials + extra))
    with pytest.raises(ValueError, match="colour-blind-safe"):
        search_report_main([str(run)])


def test_the_report_says_when_no_test_result_was_recorded_or_nothing_met_the_limit(run_dir):
    run, trials = run_dir
    for t in trials:
        t["met"] = False
    (run / "search" / "trials.json").write_text(json.dumps(trials))  # still running, so rebuilt into partial/
    search_report_main([str(run)])
    text = (run / "search" / "partial" / "search_report.md").read_text()
    assert "not recorded for this report" in text and "No setting met the limit" in text


def test_best_so_far_keeps_the_lowest_full_try_under_the_limit():
    t = lambda c, v, full=True, met=True: {"searcher": "a", "cost": c, "full": full, "met": met, "values": [v]}  # noqa: E731
    trials = [t(1, 9.0, met=False), t(1 / 3, 5.0, full=False), t(1, 8.0), t(1, 8.5), t(0, FAILED, met=False), t(1, 7.0)]
    assert best_so_far(trials, "a") == [(pytest.approx(7 / 3), 8.0), (pytest.approx(10 / 3), 8.0),
                                         (pytest.approx(10 / 3), 8.0), (pytest.approx(13 / 3), 7.0)]
    assert best_so_far(trials, "b") == []


def test_tokens_per_second_only_for_a_real_time():
    assert _tokens_per_s(4.0) == 250.0
    assert _tokens_per_s(None) is None and _tokens_per_s(0) is None and _tokens_per_s(FAILED) is None


def test_a_vllm_first_objective_gets_its_own_best_column(run_dir):
    run, _ = run_dir
    cfg = json.loads((run / "config.json").read_text())
    cfg["search"]["objectives"] = OBJ[::-1]  # "vLLM ..." keeps its capitals in the comparison key
    (run / "config.json").write_text(json.dumps(cfg))
    search_report_main([str(run)])
    text = (run / "search" / "search_report.md").read_text()
    assert "Best vLLM decode time per token (lower is better)" in text and "lowest vLLM decode time per token" in text


def _best(trials):
    met = [t for t in trials if t["full"] and t["met"]]
    return {o: min(met, key=lambda t, i=i: t["values"][i], default=None) for i, o in enumerate(OBJ)}


def test_trade_off_labels_match_the_section_1_rows_and_callouts_sit_on_the_best_points(run_dir):
    run, trials = run_dir
    search_report_main([str(run)])
    text = (run / "search" / "search_report.md").read_text()
    front = sorted(pareto(trials), key=lambda t: t["values"][0])
    labels = front_labels(front, "vllm_decode_ms_per_token", "ppl_val")
    assert len(labels) == len(front)
    for k, (tag, (ms, ppl)) in enumerate(labels, 1):
        row = next(r for r in text.splitlines() if r.startswith(f"| {k} |"))  # the same setting in section 1
        assert tag == f"#{k}" and f"| {ppl:.4g} |" in row and f"| {ms:.4g} |" in row
    best = _best(trials)
    callouts = dict(goal_callouts(best, "vllm_decode_ms_per_token", "ppl_val"))
    for o, t in best.items():
        xy = (t["metrics"]["vllm_decode_ms_per_token"], t["metrics"]["ppl_val"])
        assert callouts[f"lowest {label(o)[:1].lower() + label(o)[1:]}"] == xy


def test_a_setting_best_on_both_goals_gets_one_panel_and_one_callout_naming_both():
    t = {"metrics": {"ppl_val": 9.0, "vllm_decode_ms_per_token": 4.0}}
    best = {"ppl_val": t, "vllm_decode_ms_per_token": t}
    assert best_panels(best) == [(["ppl_val", "vllm_decode_ms_per_token"], t)]
    assert goal_callouts(best, "vllm_decode_ms_per_token", "ppl_val") == [
        ("lowest perplexity (validation half) and lowest vLLM decode time per token", (4.0, 9.0))]
    assert best_panels({"ppl_val": None}) == []


def test_measured_text_follows_the_objectives_and_skips_failed_values():
    t = {"metrics": {"ppl_val": 9.0, "ppl_heldout": 9.5, "decode_ms_per_token_mean": FAILED,
                     "predicted_weight_memory_gb": 0.8}}
    text = measured_text(t, ["ppl_val", "decode_ms_per_token_mean"], "predicted_weight_memory_gb")
    assert "Perplexity (validation half): 9" in text and "Perplexity (held-out half): 9.5" in text
    assert "0.8 GB" in text and "ms" not in text and "tokens/s =" not in text and "lower is better" in text


@pytest.mark.parametrize("keep", ["none", "one"])
def test_the_report_is_written_with_no_tries_or_a_single_one(run_dir, keep):
    run, trials = run_dir
    kept = [] if keep == "none" else [next(t for t in trials if t["met"])]
    (run / "search" / "trials.json").write_text(json.dumps(kept))
    (run / "search" / "search_report.md").write_text("old")  # written after trials.json: a finished run
    search_report_main([str(run)])
    d = run / "search"
    text = (d / "search_report.md").read_text()
    assert text.startswith("# Settings search")
    for png in re.findall(r"!\[[^]]*\]\(([^)]+)\)", text):
        assert (d / png).exists(), png
    assert (d / "best_settings.png").exists() == (keep == "one")
