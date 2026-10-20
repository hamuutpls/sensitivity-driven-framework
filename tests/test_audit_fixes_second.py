"""Fixes from the 2026-10-20 ponytail audit (21, 26, 27, 30) and follow-ups to the 2026-10-19 fixes. Every test here
fails on main before this change."""

import json
import runpy
import sys
import types
from pathlib import Path

import pytest

from sdf.config import FrameworkConfig

REPO = Path(__file__).parent.parent


def _reporter(tmp_path, main_metrics):
    from sdf.reporting.reporter import StageReporter
    from sdf.requirements import DeploymentRequirement

    return StageReporter(1, tmp_path, "t", {}, {}, {}, DeploymentRequirement(), list(main_metrics))


# 21, 30 -------------------------------------------------------------------------------------------------------
def test_backend_perplexity_is_shown_but_not_won_or_lost(tmp_path):
    import sdf.eval.llamacpp  # noqa: F401 - registers the backend columns
    import sdf.eval.trtllm_backend  # noqa: F401
    import sdf.eval.vllm_backend  # noqa: F401
    from sdf.reporting.metrics import METRICS

    backend = ["llamacpp_ppl_val", "vllm_ppl_val", "trtllm_ppl_val"]
    rep = _reporter(tmp_path, ["ppl_val", *backend])
    with rep.method("m", "original") as r:
        r.metrics.update({"ppl_val": 10.0, **dict.fromkeys(backend, 10.0)})
    with rep.method("m", "framework") as r:
        r.metrics.update({"ppl_val": 9.0, **dict.fromkeys(backend, 9.0)})
    rep.compute_deltas()
    assert rep.verdict(rep.rows[1]) == (["ppl_val"], [], [])  # perplexity counts once
    assert all("vs_original_abs" in rep.rows[1].deltas[m] for m in backend)  # still in the tables
    assert all("re-rounded" in METRICS[m].label for m in backend)


# 2026-10-19 #16 follow-up -------------------------------------------------------------------------------------
def test_excel_colours_the_held_out_half_and_one_off_costs(tmp_path):
    from openpyxl import load_workbook

    from sdf.reporting.excel import GREEN, RED, write_workbook

    rep = _reporter(tmp_path, ["ppl_val", "ppl_heldout", "build_time_s"])
    with rep.method("m", "original") as r:
        r.metrics.update({"ppl_val": 10.0, "ppl_heldout": 10.0, "build_time_s": 1.0})
    with rep.method("m", "framework") as r:
        r.metrics.update({"ppl_val": 9.0, "ppl_heldout": 9.0, "build_time_s": 5.0})
    rep.compute_deltas()
    assert rep.rows[1].verdict == {"wins": ["ppl_val"], "losses": [], "one_off": ["build_time_s"]}
    write_workbook(rep, tmp_path / "s.xlsx")
    ws = load_workbook(tmp_path / "s.xlsx")["Summary"]

    def fill(header):
        return ws.cell(row=3, column=next(c.column for c in ws[1] if c.value == header)).fill.fgColor.rgb
    assert fill("ppl_heldout Δ vs original") == GREEN.fgColor.rgb
    assert fill("build_time_s Δ vs original") == RED.fgColor.rgb


# 26 -----------------------------------------------------------------------------------------------------------
def test_the_notebook_does_not_use_google_drive():
    cells = json.loads((REPO / "notebooks" / "colab_run.ipynb").read_text())["cells"]
    code = "".join("".join(c["source"]) for c in cells if c["cell_type"] == "code")
    assert "google.colab" not in code and "drive.mount" not in code and "DRIVE" not in code
    assert 'main.OUTPUT_ROOT = f"{ROOT}/results"' in code


# 27 -----------------------------------------------------------------------------------------------------------
class _Stop(Exception):
    pass


@pytest.mark.parametrize("argv, namespace, expected", [
    (["rtn", "gptq"], {}, ["rtn", "gptq"]),  # %run colab_stage1.py rtn gptq
    ([], {"METHODS": ["awq"]}, ["awq"]),  # colab_stage1_joint.py sets METHODS and exec()s the script
])
def test_colab_stage1_runs_the_methods_it_is_given(monkeypatch, argv, namespace, expected):
    seen = {}
    fake = types.ModuleType("main")

    def run():
        seen.update(methods=fake.STAGE1_METHODS, out=fake.OUTPUT_ROOT)
        raise _Stop
    fake.main = run
    monkeypatch.setitem(sys.modules, "main", fake)
    monkeypatch.setattr(sys, "argv", ["colab_stage1.py", *argv])
    monkeypatch.chdir(REPO)  # the script changes directory to the repository
    with pytest.raises(_Stop):
        runpy.run_path(str(REPO / "scripts" / "colab_stage1.py"), init_globals=namespace)
    assert seen == {"methods": expected, "out": "/content/thesis_compression/results"}


# Stage 0 text -------------------------------------------------------------------------------------------------
def test_prune_sweep_names_the_bits_it_really_stores():
    from sdf.stage0.prune_sweep import stored_bits

    s0 = FrameworkConfig().stage0
    s0.prune_sweep_quantize = False
    assert stored_bits(s0) == (s0.baseline_bits,) * 3
    s0.prune_sweep_quantize = True
    assert stored_bits(s0) == (s0.uniform_bits, s0.protected_bits, s0.compressed_bits)


# review follow-ups --------------------------------------------------------------------------------------------
def test_a_timed_out_backend_program_takes_its_children_with_it(tmp_path):
    import subprocess
    import time

    from sdf.eval.worker import run_killing_group

    pid_file = tmp_path / "child"
    child = f"import subprocess, sys; p = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)']); " \
            f"open({str(pid_file)!r}, 'w').write(str(p.pid)); p.wait()"
    with pytest.raises(subprocess.TimeoutExpired):
        run_killing_group([sys.executable, "-c", child], 3)
    pid = int(pid_file.read_text())
    time.sleep(0.5)
    status = Path(f"/proc/{pid}/status")
    assert not status.exists() or "zombie" in status.read_text()  # killed (a zombie until init reaps it)


def test_a_row_saved_before_the_latency_mode_was_recorded_is_still_compared(tmp_path):
    rep = _reporter(tmp_path, ["decode_ms_per_token_mean"])
    with rep.method("m", "original") as r:
        r.metrics.update({"decode_ms_per_token_mean": 4.0})
    with rep.method("m", "framework") as r:
        r.metrics.update({"decode_ms_per_token_mean": 8.0, "latency_mode": "cuda_graph"})
    rep.compute_deltas()
    assert rep.verdict(rep.rows[1]) == ([], ["decode_ms_per_token_mean"], [])


def test_aqlm_does_not_carry_calibration_settings_it_does_not_use():
    from sdf.stages.methods import METHODS

    assert not METHODS["aqlm"].calibrated and METHODS["gptq"].calibrated
