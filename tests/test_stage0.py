import json
import math

import pytest
import torch
from openpyxl import load_workbook

from sdf.data import calibration_batches, eval_windows
from sdf.eval.metrics import measure_model, model_size_gb
from sdf.run import start_run
from sdf.search_space import PER_CHANNEL, SEARCH_SPACE
from sdf.stage0.activation import (ActivationPlan, ActivationProfile, activation_plan_from_weights,
                                   plan_activations, predicted_rise, profile_activations, uniform_activation_plan)
from sdf.stage0.kv_cache import KVPlan
from sdf.stage0.planner import (CompressionPlan, baseline_cost, budget_matched_plan, guarded_layers, plan_compression,
                                predict_cost, uniform_plan)
from sdf.stage0.run import _MEASURE_PLAIN, profile_key, run_stage0
from sdf.stage0.sensitivity import SCORES, SensitivityProfile, normalize, outlier_layers, profile_sensitivity
from sdf.utils.model_info import count_parameters
from conftest import fake_texts


TOY_SCORES = [0.0, 0.3, 0.5, 1.0]


def toy_profile():
    return SensitivityProfile(raw_scores=[1, 2, 3, 4], layer_numel=[1000] * 4, layer_rows=[10] * 4, other_numel=500)


def test_every_score_has_a_plain_explanation():
    # the report's "how this stage works" text looks the score up; a missing one crashed a finished hessian run
    assert set(SCORES) <= set(_MEASURE_PLAIN)


def test_normalize():
    assert normalize([2.0, 4.0, 3.0], "minmax") == [0.0, 1.0, 0.5]
    assert normalize([2.0, 9.0, 3.0]) == [0.0, 1.0, 0.5]  # rank ignores magnitudes
    assert normalize([1.0, 2.0, 2.0, 3.0]) == [0.0, 0.5, 0.5, 1.0]  # ties share their average rank
    assert normalize([5.0, 5.0]) == [0.0, 0.0] and normalize([5.0, 5.0], "minmax") == [0.0, 0.0]
    with pytest.raises(ValueError):
        normalize([1.0, 2.0], "zscore")


def test_outlier_layer_does_not_protect_everything():
    # TinyLlama on WikiText-2: layer 0 is a low outlier, the other 21 layers sit in 6596..8647.
    raw = [2705.0] + [6596.0 + i * (8647.0 - 6596.0) / 20 for i in range(21)]

    minmax = plan_compression(normalize(raw, "minmax"), 0.5, 0.3, 8, 4)
    assert len(minmax.protected_layers) == 21  # the degenerate plan seen on the real model

    plan = plan_compression(normalize(raw), 0.5, 0.3, 8, 4)
    assert len(plan.protected_layers) == 11  # threshold 0.5 protects the more sensitive half
    assert 0 in plan.compressed_layers and 21 in plan.protected_layers
    assert len(plan_compression(normalize(raw), 0.8, 0.3, 8, 4).protected_layers) == 5


def test_budget_matched_plan_fits_uniform_size():
    raw = [2705.0] + [6596.0 + i * (8647.0 - 6596.0) / 20 for i in range(21)]
    scores = normalize(raw)
    prof = SensitivityProfile(raw, [44_000_000] * 22, [12_000] * 22, 260_000_000)
    cost = lambda p: predict_cost(p, prof, 128, 32, 16)  # noqa: E731
    budget = cost(uniform_plan(scores, 4, 0.0)).weight_memory_gb
    plan = budget_matched_plan(scores, budget, 0.3, 8, 4, cost)
    assert cost(plan).weight_memory_gb <= budget
    k = len(plan.protected_layers)
    assert 0 < k < 22
    assert set(plan.protected_layers) == set(sorted(range(22), key=lambda i: raw[i])[-k:])  # the most sensitive
    one_more = budget_matched_plan(scores, budget * 10, 0.3, 8, 4, cost)
    assert len(one_more.protected_layers) == 22  # a generous budget protects everything


def test_no_prune_budget_plan_matches_size_with_bits_only():
    raw = [2705.0] + [6596.0 + i * (8647.0 - 6596.0) / 20 for i in range(21)]
    scores = normalize(raw)
    prof = SensitivityProfile(raw, [44_000_000] * 22, [12_000] * 22, 260_000_000)
    cost = lambda p: predict_cost(p, prof, 128, 32, 16)  # noqa: E731
    uni = cost(uniform_plan(scores, 4, 0.0))
    plan = budget_matched_plan(scores, uni.weight_memory_gb, 0.0, 8, 3, cost)
    c = cost(plan)
    assert c.weight_memory_gb <= uni.weight_memory_gb and c.sparsity == 0
    # 8k + 3(22 - k) <= 4 * 22 bits per weight  ->  k = 4 protected layers, the 4 most sensitive
    assert set(plan.protected_layers) == set(sorted(range(22), key=lambda i: raw[i])[-4:])
    assert {lp.bit_width for lp in plan.layers if not lp.protected} == {3}
    assert c.sensitivity_exposure < uni.sensitivity_exposure


def test_outlier_layers():
    raw = [2705.0] + [6596.0 + i * (8647.0 - 6596.0) / 20 for i in range(21)]
    assert outlier_layers(raw) == [0]
    assert outlier_layers([1.0, 1.1, 0.9, 1.05, 0.95]) == []


def test_plan_and_uniform():
    plan = plan_compression(TOY_SCORES, 0.5, 0.25, protected_bits=8, compressed_bits=4)
    assert plan.protected_layers == [2, 3] and plan.compressed_layers == [0, 1]
    assert [(lp.bit_width, lp.pruning_ratio) for lp in plan.layers] == [(4, 0.25), (4, 0.25), (8, 0.0), (8, 0.0)]
    uni = uniform_plan(TOY_SCORES, 4, 0.0)
    assert {lp.bit_width for lp in uni.layers} == {4} and not uni.protected_layers
    with pytest.raises(ValueError):
        plan_compression(TOY_SCORES, 0.5, 1.5, 8, 4)


def test_guard_blocks_pruning_of_critical_layers():
    # TinyLlama layer removal: 0, 2, 7, 21, 1 hurt most. Here: layer 0 is critical but ranks least sensitive,
    # as it does under gradient x weight.
    removal = [1190.0, 20.0, 5.0, 900.0]
    guard = guarded_layers(removal, 2)
    assert guard == {0, 3}
    assert guarded_layers(removal, 0) == frozenset()
    with pytest.raises(ValueError):
        guarded_layers(removal, -1)

    plan = plan_compression(TOY_SCORES, 0.9, 0.3, 8, 4, guard)
    assert plan.guarded_layers == [0, 3]
    assert [lp.pruning_ratio for lp in plan.layers] == [0.0, 0.3, 0.3, 0.0]
    assert plan.layers[0].bit_width == 4  # the guard blocks pruning only, bits still follow the plan

    prof = toy_profile()
    cost = lambda p: predict_cost(p, prof, 100, 32, 16)  # noqa: E731
    budget = cost(uniform_plan(TOY_SCORES, 4, 0.0)).weight_memory_gb
    same = budget_matched_plan(TOY_SCORES, budget, 0.5, 8, 4, cost, guard)
    assert all(same.layers[i].pruning_ratio == 0 for i in guard)
    assert cost(same).weight_memory_gb <= budget  # the unpruned guarded layers are paid for within the budget

    assert CompressionPlan.from_dict(plan.to_dict()) == plan


def test_activation_plan_follows_weight_plan():
    plan = plan_compression(TOY_SCORES, 0.5, 0.3, 8, 4, frozenset({0}))
    act = activation_plan_from_weights(plan, 8, 4)
    assert [lp.act_bits for lp in act.layers] == [8, 4, 8, 8]  # protected 2, 3 and guarded 0
    assert act.avg_bits == 7.0
    assert ActivationPlan.from_dict(act.to_dict()) == act


def test_measured_activation_plan_spends_bits_where_rounding_hurts():
    # rise at [4, 8] bits per layer: layers 1 and 3 suffer at 4 bits
    prof = ActivationProfile([4, 8], [[0.01, 0.0], [0.5, 0.0], [-0.02, 0.01], [0.3, 0.001]])
    plan = plan_activations(prof, 6.0)
    assert [lp.act_bits for lp in plan.layers] == [4, 8, 4, 8]
    assert plan.avg_bits == 6.0 and [lp.layer for lp in plan.layers if lp.protected] == [1, 3]
    assert predicted_rise(plan, prof) == pytest.approx(0.01 + 0.0 + 0.0 + 0.001)  # negative rise clipped to 0
    assert predicted_rise(uniform_activation_plan(4, 4), prof) == pytest.approx(0.81)
    assert predicted_rise(uniform_activation_plan(4, 6), prof) is None  # 6 bits were not measured


def test_profile_activations_restores_the_model(tiny_llama, tokenizer):
    batches = calibration_batches(fake_texts("wikitext2", "train"), tokenizer, 4, 16, 2, 0)
    before = [p.clone() for p in tiny_llama.parameters()]
    prof = profile_activations(tiny_llama, batches, [2, 8], 8, device="cpu")
    assert prof.num_layers == 4 and all(len(r) == 2 for r in prof.rise)
    assert all(math.isfinite(x) for r in prof.rise for x in r)
    assert max(abs(r[0]) for r in prof.rise) > max(abs(r[1]) for r in prof.rise)  # 2 bits change more than 8
    assert all(torch.equal(a, b) for a, b in zip(before, tiny_llama.parameters()))
    assert not any(m._forward_pre_hooks for m in tiny_llama.modules())  # hooks removed


def test_predict_cost():
    prof = toy_profile()
    fp16 = baseline_cost(prof, 16)
    assert fp16.weight_memory_gb == pytest.approx(4500 * 2 / 1e9)
    uni = predict_cost(uniform_plan(TOY_SCORES, 4, 0.0), prof, group_size=100, group_overhead_bits=32, baseline_bits=16)
    # per layer: 1000 * 4 bits + 10 groups * 32 bits = 4320 bits; + 500 * 16 bits outside the layers
    assert uni.weight_memory_gb == pytest.approx((4 * 4320 + 500 * 16) / 8 / 1e9)
    assert uni.sensitivity_exposure == pytest.approx(0.75)
    per_ch = predict_cost(uniform_plan(TOY_SCORES, 4, 0.0), prof, PER_CHANNEL, 32, 16)
    assert per_ch.per_layer_mb[0] == pytest.approx((4000 + 10 * 32) / 8 / 1e6)
    fw = predict_cost(plan_compression(TOY_SCORES, 0.5, 0.3, 8, 4), prof, 100, 32, 16)
    assert fw.sensitivity_exposure < uni.sensitivity_exposure  # sensitive layers are spared
    assert fw.sparsity == pytest.approx(2 * 300 / 4500)


def test_profile_sensitivity_on_tiny_llama(tiny_llama):
    frozen = tiny_llama.model.embed_tokens.weight
    frozen.requires_grad_(False)
    batches = [torch.randint(0, 64, (2, 16), generator=torch.Generator().manual_seed(i)) for i in range(3)]
    prof = profile_sensitivity(tiny_llama, batches, device="cpu")
    assert prof.num_layers == 4 and all(s > 0 for s in prof.raw_scores)
    assert prof.cost["calibration_tokens"] == 96
    # q, o: 32x32; k, v: 32x16 (2 KV heads); gate, up, down: 32x64
    assert prof.layer_numel == [2 * 1024 + 2 * 512 + 3 * 2048] * 4
    assert all(p.grad is None for p in tiny_llama.parameters())
    assert frozen.requires_grad is False
    assert SensitivityProfile.from_dict(prof.to_dict()) == prof


def test_data_windows(tokenizer):
    a = calibration_batches(fake_texts("x", "train"), tokenizer, n_samples=6, seq_len=8, batch_size=4, seed=1)
    b = calibration_batches(fake_texts("x", "train"), tokenizer, n_samples=6, seq_len=8, batch_size=4, seed=1)
    assert [t.shape for t in a] == [(4, 8), (2, 8)] and all(torch.equal(x, y) for x, y in zip(a, b))
    val, held = eval_windows(fake_texts("x", "test"), tokenizer, seq_len=16, max_windows=5)
    assert val.shape == (2, 16) and held.shape == (3, 16)


def test_measure_model(tiny_llama, tokenizer, small_cfg):
    val, held = eval_windows(fake_texts("x", "test"), tokenizer, 16, 4)
    metrics, raw = measure_model(tiny_llama, val, held, small_cfg.eval, torch.device("cpu"))
    assert metrics["ppl_val"] > 1 and metrics["model_size_gb"] == pytest.approx(model_size_gb(tiny_llama))
    assert sum(r["measurement"] == "latency" for r in raw) == small_cfg.eval.latency_repeats


def test_run_stage0_end_to_end(tiny_llama, tokenizer, small_cfg, monkeypatch):
    ctx = start_run(small_cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16, "sensitive_threshold": 0.5})
    res = run_stage0(ctx, cand, model=tiny_llama, tokenizer=tokenizer, text_loader=fake_texts)

    stage_dir = ctx.run_dir / "stage_0"
    for name in ("report.md", "stage_0_comparison.xlsx", "results.json", "compression_plan.json",
                 "compression_plan_budget_matched_no_prune.json", "sensitivity_profile.json",
                 "kv_cache_plan.json", "kv_cache_plan_bits_only.json", "kv_profile.json", "activation_plan.json",
                 "activation_profile.json"):
        assert (stage_dir / name).exists(), name
    data = json.loads((stage_dir / "results.json").read_text())
    rows = {f"{r['method']}/{r['variant']}": r for r in data["rows"]}
    assert set(rows) == {"baseline/fp16", "allocation/original", "allocation/framework",
                         "allocation_same_size/framework", "allocation_same_size_no_prune/framework",
                         "kv_cache/original", "kv_cache/framework", "kv_cache_bits_only/framework",
                         "activations/original", "activations/framework", "activations_from_weights/framework"}
    kv_fw, kv_un = rows["kv_cache/framework"]["metrics"], rows["kv_cache/original"]["metrics"]
    assert kv_fw["predicted_kv_memory_gb"] <= kv_un["predicted_kv_memory_gb"] * (1 + 1e-9)
    assert kv_fw["avg_kv_bits"] <= 4 + 1e-9
    assert rows["baseline/fp16"]["metrics"]["predicted_kv_memory_gb"] > kv_un["predicted_kv_memory_gb"]
    assert "vs_original_abs" in rows["kv_cache_bits_only/framework"]["deltas"]["predicted_kv_ppl_rise"]
    no_prune = rows["allocation_same_size_no_prune/framework"]
    assert no_prune["metrics"]["sparsity"] == 0
    assert no_prune["metrics"]["predicted_weight_memory_gb"] <= rows["allocation/original"]["metrics"][
        "predicted_weight_memory_gb"] * (1 + 1e-9)
    same = rows["allocation_same_size/framework"]
    assert same["metrics"]["predicted_weight_memory_gb"] <= rows["allocation/original"]["metrics"][
        "predicted_weight_memory_gb"] * (1 + 1e-9)
    assert "vs_original_abs" in same["deltas"]["predicted_weight_memory_gb"]  # compared to the uniform plan
    assert all(r["status"] == "ok" for r in rows.values())
    assert rows["baseline/fp16"]["metrics"]["ppl_val"] > 1
    fw = rows["allocation/framework"]
    assert fw["metrics"]["protected_layers"] == len(res.plan.protected_layers)
    assert "vs_original_abs" in fw["deltas"]["sensitivity_exposure"]

    # the plans later stages read: returned, and loadable from the JSON files
    assert len(res.plan.guarded_layers) == 1
    for p in (res.plan, res.budget_plan, res.no_prune_plan):
        assert all(p.layers[i].pruning_ratio == 0 for i in p.guarded_layers)
    assert CompressionPlan.load(stage_dir / "compression_plan.json") == res.plan
    assert CompressionPlan.load(stage_dir / "compression_plan_budget_matched.json") == res.budget_plan
    assert KVPlan.load(stage_dir / "kv_cache_plan.json") == res.kv_plan
    assert KVPlan.load(stage_dir / "kv_cache_plan_bits_only.json") == res.kv_plan_bits_only
    assert ActivationPlan.load(stage_dir / "activation_plan.json") == res.activation_plan
    assert rows["activations/framework"]["metrics"]["avg_activation_bits"] == res.activation_plan.avg_bits == 6
    assert res.activation_plan.kind == "measured" and res.activation_profile is not None
    assert rows["activations/original"]["metrics"]["avg_activation_bits"] == 8
    act_fw, act_un = rows["activations/framework"]["metrics"], rows["activations/original"]["metrics"]
    assert act_fw["predicted_act_ppl_rise"] >= act_un["predicted_act_ppl_rise"]  # fewer bits cannot help
    assert "predicted_act_ppl_rise" in rows["activations_from_weights/framework"]["metrics"]
    assert all(r["requirement"]["targets_set"] is False for r in rows.values())  # no targets in small_cfg
    assert len(data["per_layer"]) == 4
    wb = load_workbook(stage_dir / "stage_0_comparison.xlsx")
    assert wb["Per-layer"].max_row == 5
    n_params = count_parameters(tiny_llama)
    assert data["original_model"]["num_parameters"] == n_params
    assert data["original_model"]["num_layers"] == 4
    assert "original_model.num_key_value_heads" in [c.value for c in wb["Config"]["A"]]

    # when everything is cached the weights are never loaded: the config file and the profile give the facts
    import transformers

    monkeypatch.setattr(transformers.AutoConfig, "from_pretrained", lambda name: tiny_llama.config)
    # second trial with other Stage 0 hyperparameters: profile and FP16 baseline come from cache
    res2 = run_stage0(ctx, SEARCH_SPACE.make({"calib_samples": 16, "sensitive_threshold": 0.9}),
                      model=None, tokenizer=tokenizer, text_loader=fake_texts)
    assert res2.profile == res.profile
    data2 = json.loads((stage_dir / "results.json").read_text())
    assert data2["rows"][0]["info"]["cached"] is True
    assert data2["rows"][2]["info"]["profile_cached"] is True
    report = (stage_dir / "report.md").read_text(encoding="utf-8")
    assert "also includes a budget plan" in report and "Size floor" in report
    assert "## Sensitivity of every layer" in report
    assert "Unused budget: budget plan without pruning" in report
    assert "## KV cache plan" in report and "Key bits (cache)" in report and "short-term memory" in report
    assert "Share removed (budget plan)" in report and "pruning caveat" in report
    assert "Perplexity (held-out half)" in report.split("# Technical details")[1]
    assert report.split("## Summary")[1].split("##")[0].count("\n\n") <= 2  # one paragraph
    assert data2["original_model"]["num_parameters"] == n_params
    layer_part = report.split("## Sensitivity of every layer")[1].split("\n## ")[0]
    header = next(line for line in layer_part.splitlines() if line.startswith("| Layer"))
    notes = layer_part.split("**What each column means**")[1]
    assert notes.count("\n- **") == header.count("|") - 1  # one explanation per column
    assert "- **Raw score**: How much the prediction error" in layer_part
    assert "## Original model" in report and f"{n_params:,}" in report
    assert "Pruning guard" in report and "Activation plan (for Stage 2)" in report
    handoff = (stage_dir / "handoff.md").read_text(encoding="utf-8")
    for heading in ("## Original model", "## At a glance", "## Stage 1: weights", "## Stage 2: activations",
                    "## Stage 3: KV cache", "## Stage 4: evaluation", "## Search", "## Caveats"):
        assert heading in handoff, heading
    tables = [b for b in handoff.split("\n\n") if b.startswith("| ")]
    notes = [b for b in handoff.split("**What each column means**")[1:]]
    assert len(notes) == len(tables) - 1  # all but the Original model table carry column notes
    findings = report.split("## Findings")[1].split("\n## ")[0]
    assert "paid once" in findings and "time to prepare" not in findings.split("It costs more")[0]  # not a verdict
    results_rows = [l for l in report.split("## Results")[-1].splitlines() if l.startswith("| `")]
    assert results_rows and all(l.endswith("| no targets set |") for l in results_rows)


def test_guard_uses_layer_removal_when_another_score_picks_bits(tiny_llama, tokenizer, small_cfg):
    cfg = small_cfg.with_overrides({"stage0.score": "grad_x_weight", "stage0.kv_cache": False})
    ctx = start_run(cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16})
    res = run_stage0(ctx, cand, model=tiny_llama, tokenizer=tokenizer, text_loader=fake_texts, measure_fp16=False)
    assert res.profile.method == "grad_x_weight"
    assert len(res.plan.guarded_layers) == 1
    # the removal profile behind the guard was measured and cached next to the gradient one
    assert ctx.cache.path("sensitivity_profile", profile_key(ctx, cand, "layer_removal")).exists()


def test_sweep_end_to_end(tiny_llama, tokenizer, small_cfg):
    from sdf.stage0.sweep import run_sweep

    ctx = start_run(small_cfg)
    grid = {"sensitive_threshold": [0.3, 0.7], "prune_ratio_aggressive": [0.0, 0.3], "gptq_groupsize": [32, PER_CHANNEL],
            "calib_dataset": ["wikitext2", "c4"], "calib_samples": [8, 16]}
    out = run_sweep(ctx, grid, model=tiny_llama, tokenizer=tokenizer, text_loader=fake_texts)
    data = json.loads(out["json"].read_text())
    assert len(data["rows"]) == 2 * 2 * 2 * 2 * 2
    assert len(data["calibration"]) == 4 and not data["failures"]
    assert all(r["same_gb"] <= r["uniform_gb"] * (1 + 1e-9) and r["noprune_gb"] <= r["uniform_gb"] * (1 + 1e-9)
               for r in data["rows"])
    report = out["report"].read_text(encoding="utf-8")
    for heading in ("## Summary", "## 1. Threshold", "## 2. Prune ratio", "## 3. Group size", "## 4. Calibration"):
        assert heading in report
    assert load_workbook(out["xlsx"])["All plans"].max_row == len(data["rows"]) + 1


def test_ablation_scores_restore_the_model(tiny_llama):
    from sdf.stage0.sensitivity import find_decoder_layers, profile_by_ablation, skip_layer

    batches = [torch.randint(0, 64, (2, 16), generator=torch.Generator().manual_seed(i)) for i in range(2)]
    before = {k: v.clone() for k, v in tiny_llama.state_dict().items()}
    ids = batches[0]
    with torch.no_grad():
        ref = tiny_llama(input_ids=ids).logits
        with skip_layer(find_decoder_layers(tiny_llama)[1]):
            skipped = tiny_llama(input_ids=ids).logits
        after = tiny_llama(input_ids=ids).logits
    assert not torch.allclose(ref, skipped) and torch.allclose(ref, after)  # hook removed afterwards

    for method in ("layer_removal", "layer_quant"):
        prof = profile_by_ablation(tiny_llama, batches, method, device="cpu", bits=2, group_size=16)
        assert prof.method == method and prof.num_layers == 4
        assert prof.meta["baseline_ppl"] > 1 and len(prof.meta["layer_ppl"]) == 4
        assert any(s != 0 for s in prof.raw_scores)
        assert all(torch.equal(v, before[k]) for k, v in tiny_llama.state_dict().items())  # weights restored


def test_compare_scores_end_to_end(tiny_llama, tokenizer, small_cfg):
    from sdf.stage0.compare import compare_scores

    ctx = start_run(small_cfg)
    cand = SEARCH_SPACE.make({"calib_samples": 16, "gptq_groupsize": 32})
    out = compare_scores(ctx, cand, model=tiny_llama, tokenizer=tokenizer, text_loader=fake_texts)
    data = json.loads(out["json"].read_text())
    assert data["scores"] == ["grad_x_weight", "layer_removal", "layer_quant", "fisher", "taylor_ema", "hessian",
                              "movement"] and not data["failures"]
    assert all(data["agreement"][s][s] == pytest.approx(1.0) for s in data["scores"])
    assert len(data["per_layer"]) == 4
    report = out["report"].read_text(encoding="utf-8")
    assert "## Agreement between the ways" in report and "remove the layer" in report
    assert "## How well each way matches the measured damage" in report and "Hessian" in report

    # a normal Stage 0 run with the removal score
    ctx2 = start_run(small_cfg.with_overrides({"stage0.score": "layer_removal", "run.run_id": "removal"}))
    res = run_stage0(ctx2, cand, model=tiny_llama, tokenizer=tokenizer, text_loader=fake_texts, measure_fp16=False)
    assert res.profile.method == "layer_removal"
    assert "switched off" in (ctx2.run_dir / "stage_0" / "report.md").read_text(encoding="utf-8")


@pytest.mark.parametrize("method", ["fisher", "taylor_ema", "hessian", "movement"])
def test_gradient_scores_restore_the_model(tiny_llama, method):
    torch.manual_seed(0)
    batches = [torch.randint(0, 64, (2, 16)) for _ in range(3)]
    before = {n: p.clone() for n, p in tiny_llama.named_parameters()}
    prof = profile_sensitivity(tiny_llama, batches, method=method, movement_lr=1e-2)
    assert prof.method == method and len(prof.raw_scores) == 4
    assert all(math.isfinite(s) for s in prof.raw_scores) and len(set(prof.raw_scores)) > 1
    assert all(s >= 0 for s in prof.raw_scores)
    assert all(torch.equal(p, before[n]) for n, p in tiny_llama.named_parameters())
    assert all(p.grad is None for p in tiny_llama.parameters())


def test_hessian_score_matches_exact_diagonal():
    # one Linear, loss = 0.5 * sum((W x)^2) / n: the diagonal of the Hessian is known exactly, so the probe
    # estimate (averaged over many batches with different noise) should land close to 0.5 * sum(diag(H) w^2).
    from sdf.stage0.sensitivity import profile_sensitivity as prof_fn

    class Toy(torch.nn.Module):
        def __init__(self):
            super().__init__()
            self.model = torch.nn.Module()
            self.model.layers = torch.nn.ModuleList([torch.nn.Linear(4, 3, bias=False)])

        def forward(self, input_ids, labels=None):
            x = self.x
            out = self.model.layers[0](x)
            return type("O", (), {"loss": 0.5 * out.pow(2).sum() / x.shape[0]})

    torch.manual_seed(0)
    toy = Toy().double()
    toy.x = torch.randn(8, 4, dtype=torch.float64)
    w = toy.model.layers[0].weight.detach()
    diag = (toy.x.pow(2).mean(0)).expand_as(w)  # d2L/dw_ij^2 = mean_n x_nj^2
    exact = 0.5 * (diag * w * w).sum().item()
    est = prof_fn(toy, [torch.zeros(1, 1, dtype=torch.long)] * 50, method="hessian").raw_scores[0]
    assert est == pytest.approx(exact, rel=0.1)
