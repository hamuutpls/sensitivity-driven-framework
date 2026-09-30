import json

from openpyxl import load_workbook

from sdf.reporting.reporter import StageReporter
from sdf.requirements import DeploymentRequirement


def make_reporter(tmp_path):
    return StageReporter(stage=1, run_dir=tmp_path, title="Weight compression", config={"a": {"b": 1}},
                         environment={"gpu": "none"}, conditions={"seed": 0},
                         requirement=DeploymentRequirement(target_ppl=11.0), main_metrics=["ppl_val", "model_size_gb"])


def test_deltas_wins_failures_and_outputs(tmp_path):
    rep = make_reporter(tmp_path)
    with rep.method("baseline", "fp16") as r:
        r.metrics.update(ppl_val=10.0, model_size_gb=2.0)
    with rep.method("gptq", "original") as r:
        r.metrics.update(ppl_val=12.0, model_size_gb=0.6)
    with rep.method("gptq", "framework") as r:
        r.metrics.update(ppl_val=11.0, model_size_gb=0.7)
    with rep.method("awq", "framework") as r:
        raise RuntimeError("kernel missing")
    rep.add_raw("gptq", "framework", [{"measurement": "latency", "repeat": 0, "ms": 3.0}])
    rep.per_layer.append({"layer": 0, "sensitivity": 1.0})
    outputs = rep.finalize()

    fw = rep.rows[2]
    assert fw.deltas["ppl_val"]["vs_original_abs"] == -1.0
    assert fw.deltas["model_size_gb"]["vs_fp16_pct"] == -65.0
    assert rep.rows[3].status == "failed" and "kernel missing" in rep.rows[3].error
    assert rep.rows[1].requirement["met"] is False  # ppl 12 > target 11
    assert any("trades off" in f for f in rep.all_findings)

    data = json.loads(outputs["json"].read_text())
    assert [r["status"] for r in data["rows"]] == ["ok", "ok", "ok", "failed"]
    wb = load_workbook(outputs["xlsx"])
    assert wb.sheetnames == ["Summary", "Config", "Per-layer", "Raw", "Charts"]
    assert wb["Summary"].max_row == 5
    assert len(wb["Summary"].conditional_formatting) > 0
    assert wb["Charts"]._charts
    md = outputs["report"].read_text()
    assert "## Key findings" in md and "awq/framework failed" in md

    rep.finalize()  # idempotent: findings are not duplicated
    assert json.loads(outputs["json"].read_text())["findings"] == data["findings"]


def test_results_json_is_written_after_each_row(tmp_path):
    rep = make_reporter(tmp_path)
    with rep.method("baseline", "fp16") as r:
        r.metrics["ppl_val"] = 10.0
    assert len(json.loads(rep.json_path.read_text())["rows"]) == 1


def test_text_outputs_are_utf8(tmp_path):
    from sdf.utils.cache import atomic_write_text

    atomic_write_text(tmp_path / "r.md", "Stage 0 — Δ ≥ 0.5")
    assert (tmp_path / "r.md").read_bytes().decode("utf-8") == "Stage 0 — Δ ≥ 0.5"


def test_report_has_plain_language_part(tmp_path):
    rep = make_reporter(tmp_path)
    rep.plain_intro = "This stage squeezes the model."
    rep.plain_why.append("It kept the fragile parts intact.")
    with rep.method("baseline", "fp16") as r:
        r.metrics.update(ppl_val=10.0, model_size_gb=2.0)
    with rep.method("gptq", "original") as r:
        r.metrics.update(ppl_val=12.0, model_size_gb=0.6)
    with rep.method("gptq", "framework") as r:
        r.metrics.update(ppl_val=11.0, model_size_gb=0.7)
    md = rep.finalize()["report"].read_text(encoding="utf-8")

    plain, technical = md.split("# Technical details")
    for heading in ("## Summary", "## Key terms", "## Results", "## Findings",
                    "## How this stage works"):
        assert heading in plain
    assert plain.index("## Summary") < plain.index("## Key terms") < plain.index("## Results") \
        < plain.index("## Findings")
    assert "is a trade-off against the standard method" in plain
    assert "Perplexity measures how well the model predicts" in plain  # metric defined
    assert "| Version | Prediction error on test text (validation half) (lower is better) |" in plain
    assert "It kept the fragile parts intact." in plain
    assert "## Results" in technical and "`gptq/framework`" in technical  # technical tables kept
