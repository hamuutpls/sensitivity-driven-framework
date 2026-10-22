"""The settings search (sdf.search): the space, one chained evaluation, the searchers and the report."""

import copy
import json
import random
from pathlib import Path

import openpyxl
import pytest
import torch

from sdf import search
from sdf.eval import vllm_backend
from sdf.run import start_run
from sdf.search import (FAILED, Recorder, _decode, _rank_key, _sample, best_per_objective, evaluate, finalists,
                        hypervolume, pareto, run_mfbo, run_mobo, run_nsga3, run_search, technique_space, violation)
from sdf.search_space import SEARCH_SPACE
from sdf.stage0.planner import require_bits_only
from sdf.stages import runner
from sdf.stages.runner import stage1_row
from conftest import fake_texts

HF = ["ppl_val", "decode_ms_per_token_mean"]  # no vLLM in tests


@pytest.fixture
def cfg(small_cfg):
    return small_cfg.with_overrides({"requirement.target_memory_gb": 1.0, "search.objectives": HF,
                                     "search.stage1_methods": ["rtn_with_rtn_act"],
                                     "search.stage2_methods": ["unstructured_prune", "structured_prune"],
                                     "search.stage3_methods": ["rtn_kv", "h2o"]})


def _kw(tiny_llama, tokenizer):
    return dict(model_factory=lambda: copy.deepcopy(tiny_llama), tokenizer=tokenizer, text_loader=fake_texts)


def _point(cfg, **over):
    return technique_space(cfg).make({"calib_samples": 16, "prune_ratio_aggressive": 0.4, **over})


def test_technique_space_adds_a_technique_per_stage(cfg):
    space = technique_space(cfg)
    assert [p.name for p in space.params][: len(SEARCH_SPACE.params)] == [p.name for p in SEARCH_SPACE.params]
    p = _point(cfg)
    assert p["stage2_method"] == "unstructured_prune" and "stage1_plan" not in p  # Stage 1 always follows the joint plan
    with pytest.raises(ValueError, match="stage3_method"):
        space.validate({**p, "stage3_method": "gptq"})
    rng = random.Random(0)
    for _ in range(20):
        space.validate(_sample(space, rng))
    # [0, 1] vectors decode to valid points; the ends map to the first and last choice or the range ends
    lo, hi = _decode(space, [0.0] * len(space.params)), _decode(space, [1.0] * len(space.params))
    assert space.validate(lo) and space.validate(hi)
    assert lo["stage3_method"] == "rtn_kv" and hi["stage3_method"] == "h2o"
    assert lo["sensitive_threshold"] == 0.1 and hi["sensitive_threshold"] == 0.9


def _t(values, met=True, full=True, point=None, searcher="a"):
    return {"values": values, "met": met, "full": full, "point": point or {"v": values}, "searcher": searcher,
            "cost": 1.0}


def test_pareto_and_hypervolume_use_only_met_full_trials():
    ts = [_t([10, 5]), _t([12, 3]), _t([11, 6]),  # [11, 6] is beaten by [10, 5]
          _t([1, 1], met=False), _t([2, 2], full=False), _t([10, 5], point={"v": [10, 5]}, searcher="b")]
    assert [t["values"] for t in pareto(ts)] == [[10, 5], [12, 3]]  # the repeat of [10, 5] once
    # area dominated up to (20, 10): [10..20]x[5..10] + [12..20]x[3..5] = 50 + 16
    assert hypervolume(ts, [20, 10]) == pytest.approx(66)
    assert hypervolume([_t([1, 1], met=False)], [20, 10]) == 0
    assert violation({"met": False, "shortfall": {"target_memory_gb": 0.25}}) == 0.25
    assert violation({"met": True, "shortfall": {}}) == 0


def test_finalists_are_the_pareto_trials_and_the_best_per_objective_once_each():
    ts = [_t([10, 5]), _t([10, 4]), _t([12, 3]), _t([1, 1], met=False), _t([9, 9], full=False),
          _t([12, 3], point={"v": [12, 3]}, searcher="b")]
    best = best_per_objective(ts, ["ppl_val", "ms"])
    assert best["ppl_val"]["values"] == [10, 4]  # the tie on the first objective goes to the better second one
    assert best["ms"]["values"] == [12, 3] and best["ms"]["searcher"] == "a"
    assert [t["values"] for t in finalists(ts, ["ppl_val", "ms"])] == [[10, 4], [12, 3]]
    assert best_per_objective([_t([1, 1], met=False)], ["ppl_val"]) == {"ppl_val": None}
    assert finalists([_t([1, 1], met=False)], ["ppl_val"]) == []


def test_evaluate_builds_the_chained_model_and_caches_it(cfg, tiny_llama, tokenizer):
    ctx = start_run(cfg)
    kw = _kw(tiny_llama, tokenizer)
    r = evaluate(ctx, _point(cfg, stage2_method="structured_prune", stage3_method="h2o"), **kw)
    assert "error" not in r, r
    m = r["metrics"]
    assert m["ppl_val"] > 1 and m["avg_activation_bits"] < 16 and m["decoder_weights_removed"] > 0
    assert m["predicted_weight_memory_gb"] < 1.0 and r["met"]  # the toy model is far below 1 GB
    rows = json.loads((ctx.run_dir / "search" / "points" / r["dir"].rsplit("/", 1)[1] / "stage_3" /
                       "results.json").read_text())["rows"]
    assert {x["variant"] for x in rows} == {"fp16", "framework"}  # Our method only
    assert next(x for x in rows if x["variant"] == "framework")["info"]["built_on"] == "structured_prune"
    # cached: the same point is not built again, a different one is
    again = evaluate(ctx, _point(cfg, stage2_method="structured_prune", stage3_method="h2o"), **kw)
    assert again["cached"] and not r["cached"] and {**again, "cached": False} == r
    other = evaluate(ctx, _point(cfg, stage3_method="rtn_kv"), **kw)
    assert other["dir"] != r["dir"]
    # fewer windows (MFBO's rough rung) is its own entry
    rough = evaluate(ctx, _point(cfg, stage2_method="structured_prune", stage3_method="h2o"), max_windows=2, **kw)
    assert rough["dir"] != r["dir"] and rough["metrics"]["ppl_val"] != m["ppl_val"]


def test_a_tight_memory_limit_is_not_met_and_says_by_how_much(cfg, tiny_llama, tokenizer, monkeypatch):
    rows = []
    monkeypatch.setattr(runner, "stage1_row", lambda name: rows.append(name) or stage1_row(name))
    ctx = start_run(cfg.with_overrides({"requirement.target_memory_gb": 1e-9}))
    r = evaluate(ctx, _point(cfg), **_kw(tiny_llama, tokenizer))
    assert rows[-1] == "rtn_with_rtn_act_joint"  # the built model follows the joint weight x activation plan
    assert not r["met"] and r["shortfall"]["target_memory_gb"] > 0
    # the size follows from the plan, so the model is not built: no Stage 3 report, no perplexity
    assert r["not_built"] and "ppl_val" not in r["metrics"] and not (Path(r["dir"]) / "stage_3").exists()
    t = Recorder(ctx, "x", 6, **_kw(tiny_llama, tokenizer))(_point(cfg))
    assert not t["ok"] and 0 < t["violation"] < FAILED and "not built" in t["error"]
    assert t["cost"] == 0  # only built models use up the budget


def test_tries_over_the_limit_do_not_use_up_the_budget_but_are_capped(cfg, tiny_llama, tokenizer):
    c = cfg.with_overrides({"requirement.target_memory_gb": 1e-9, "search.max_tries_per_budget": 3})
    rec = Recorder(start_run(c), "mfbo", 6, **_kw(tiny_llama, tokenizer))
    run_mfbo(rec, technique_space(c), budget=1, seed=0, eta=2, rungs=2)
    assert len(rec.trials) == 3 and rec.spent == 0  # nothing built, so it stops at budget x 3 tries


def test_vllm_latency_is_measured_on_the_stage2_model(cfg, tiny_llama, tokenizer, monkeypatch):
    seen = []

    def measure(model, tok, plan, val, held, ev, base, group, acts=None, act_group_size=128):
        if plan is not None:
            require_bits_only(plan, "the export")  # as the real exports do
        ws = [m.weight for m in model.model.layers.modules() if isinstance(m, torch.nn.Linear)]
        seen.append((plan, acts, sum((w == 0).sum().item() for w in ws) / sum(w.numel() for w in ws)))
        return {"vllm_decode_ms_per_token": 3.0, "vllm_ppl_val": 9.0, "vllm_size_gb": 0.1}

    monkeypatch.setattr(vllm_backend, "measure", measure)
    c = cfg.with_overrides({"eval.vllm_python": "py", "search.objectives": ["ppl_val", "vllm_decode_ms_per_token"]})
    ctx = start_run(c)
    rec = Recorder(ctx, "x", 6, **_kw(tiny_llama, tokenizer))
    t = rec(_point(c))
    chained = [x for x in seen if x[0] is not None]  # the other call is the uncompressed row's (cached)
    assert t["ok"] and t["values"][1] == 3.0 and len(chained) == 1, t["error"]
    plan, acts, zero_share = chained[0]
    row = next(x for x in json.loads((ctx.run_dir / "search" / "points" / t["dir"].rsplit("/", 1)[1] / "stage_3" /
                                      "results.json").read_text())["rows"] if x["variant"] == "framework")
    # the model is pruned (at least the planned share is zero), the export gets the bits only, activations rounded
    assert zero_share >= row["metrics"]["decoder_weights_removed"] > 0.1
    assert max(lp.pruning_ratio for lp in plan.layers) == 0 and acts is not None
    assert "Stage 2 model" in row["info"]["backends"] and "vllm_error" not in row["info"]

    # a plan vLLM cannot load: the try fails and says why
    def refuse(*a, **k):
        raise vllm_backend.VLLMError("3-bit GPTQ is not supported")

    monkeypatch.setattr(vllm_backend, "measure", refuse)
    t = rec(_point(c, prune_ratio_aggressive=0.5))
    assert not t["ok"] and t["values"] == [FAILED, FAILED] and "3-bit GPTQ" in t["error"]
    assert not rec(_point(c, prune_ratio_aggressive=0.5))["cached"]  # a failure is tried again, not cached


def test_a_vllm_objective_without_vllm_stops_before_any_try(cfg):
    c = cfg.with_overrides({"search.objectives": ["ppl_val", "vllm_decode_ms_per_token"]})
    with pytest.raises(ValueError, match="eval.vllm_python"):
        run_search(start_run(c))


def test_mfbo_measures_the_best_on_more_windows(cfg, tiny_llama, tokenizer):
    rec = Recorder(start_run(cfg), "mfbo", 6, **_kw(tiny_llama, tokenizer))
    run_mfbo(rec, technique_space(cfg), budget=2, seed=0, eta=2, rungs=2)
    rungs = [t["windows"] for t in rec.trials]
    assert rungs[:2] == [3, 3] and rungs[2] == 6  # 2 rough tries, the better one on all the windows
    rough = rec.trials[:2]
    assert rec.trials[2]["point"] == min(rough, key=lambda t: _rank_key(t, rough))["point"]  # the better rough one
    assert rec.spent >= 2


def test_run_search_compares_the_searchers_and_writes_the_report(cfg, tiny_llama, tokenizer):
    c = cfg.with_overrides({"search.budget": 2, "search.mfbo_eta": 3})
    ctx = start_run(c)
    out = run_search(ctx, **_kw(tiny_llama, tokenizer))
    trials = json.loads(out["search_trials"].read_text())
    assert {t["searcher"] for t in trials} == {"mobo", "mfbo", "nsga3"}
    for name in ("mobo", "nsga3", "mfbo"):
        assert sum(t["cost"] for t in trials if t["searcher"] == name) >= 2
    wb = openpyxl.load_workbook(out["search_xlsx"])
    assert wb.sheetnames == ["Trials", "Pareto", "Searcher comparison", "Best configs"]
    assert wb["Trials"].max_row == len(trials) + 1
    text = out["search_report"].read_text()
    assert "Best trade-off settings" in text and "Search methods compared" in text and "at most 1.0" in text
    assert out["search_hypervolume"].exists()


class _AllOverTheLimit(Recorder):
    """A Recorder whose every point is over the limit (not built, cost 0), without building anything."""

    def __init__(self, cfg):
        self.searcher, self.full, self.trials = "nsga3", 6, []
        self.objectives, self.tries_per_budget = cfg.search.objectives, 3

    def __call__(self, point, windows=None):
        t = {"cost": 0.0, "ok": False, "met": False, "values": [FAILED] * len(self.objectives), "violation": 1.0,
             "point": point}
        self.trials.append(t)
        return t


def test_nsga3_keeps_asking_while_tries_are_not_built_and_stops_at_the_cap(cfg):
    rec = _AllOverTheLimit(cfg)
    run_nsga3(rec, technique_space(cfg), budget=2, seed=0, pop_size=4)
    assert len(rec.trials) == 6 and rec.spent == 0  # budget x 3 tries, none built


def _fake_evaluate(over=lambda p: p["sensitive_threshold"] > 0.3, crash=lambda p: False):
    """evaluate() without building: a point is over the limit by `over`, crashes by `crash` (before building, or
    while building when it returns "building"), else measured from its settings (lower threshold, lower perplexity)."""
    seen = set()

    def fake(ctx, point, max_windows=None, **kw):
        key = json.dumps([point, max_windows], sort_keys=True)
        base = {"dir": "d", "cached": key in seen}
        if c := crash(point):  # a failure is not cached
            return {**base, "error": "RuntimeError: boom", **({"attempted": True} if c == "building" else {})}
        size = 0.5 + point["sensitive_threshold"]
        seen.add(key)
        if over(point):
            return {**base, "metrics": {"predicted_weight_memory_gb": size}, "errors": {},
                    "not_built": "over the limit, so not built", "met": False,
                    "shortfall": {"target_memory_gb": size - 0.8}}
        return {**base, "metrics": {"ppl_val": 10 + point["sensitive_threshold"] + (max_windows or 0) / 100,
                                    "decode_ms_per_token_mean": 1 + point["prune_ratio_aggressive"],
                                    "predicted_weight_memory_gb": size}, "errors": {}, "met": True, "shortfall": {}}
    return fake


def test_only_models_built_for_this_searcher_use_up_the_budget(cfg, monkeypatch):
    crash = {128: "before", 64: "building"}
    monkeypatch.setattr(search, "evaluate", _fake_evaluate(crash=lambda p: crash.get(p["calib_samples"])))
    rec = Recorder(start_run(cfg), "x", 6)
    built = _point(cfg, sensitive_threshold=0.2)
    assert rec(built)["cost"] == 1 and rec(built, 3)["cost"] == 0.5
    assert rec(built)["cost"] == 0 and rec(built)["built"]  # this searcher already has it on all windows
    assert rec(_point(cfg, sensitive_threshold=0.8))["cost"] == 0  # over the limit
    t = rec(_point(cfg, sensitive_threshold=0.2, calib_samples=128))
    assert t["cost"] == 0 and not t["built"] and not t["over_limit"]  # a crash before building builds nothing
    t = rec(_point(cfg, sensitive_threshold=0.2, calib_samples=64))
    assert t["cost"] == 1 and t["built"] and not t["ok"]  # a model was built, then measuring it failed
    assert rec(_point(cfg, sensitive_threshold=0.2, calib_samples=64))["cost"] == 1  # not cached, so rebuilt
    assert rec.spent == 3.5


def test_mfbo_never_gives_more_windows_to_a_point_that_was_not_built(cfg, monkeypatch):
    monkeypatch.setattr(search, "evaluate", _fake_evaluate())
    rec = Recorder(start_run(cfg), "mfbo", 9)
    run_mfbo(rec, technique_space(cfg), budget=3, seed=0, eta=3, rungs=3)
    for t in rec.trials:
        if t["windows"] > 2:  # promoted: measured and fine on the rung below
            below = [o for o in rec.trials[: t["n"] - 1] if o["point"] == t["point"] and o["windows"] < t["windows"]]
            assert below and all(o["ok"] for o in below)
    assert any(t["windows"] == 9 for t in rec.trials) and rec.spent >= 3


def test_mobo_learns_from_points_over_the_limit(cfg, monkeypatch):
    from optuna.trial import TrialState

    monkeypatch.setattr(search, "evaluate", _fake_evaluate())
    rec = Recorder(start_run(cfg.with_overrides({"search.max_tries_per_budget": 4})), "mobo", 6)
    study = run_mobo(rec, technique_space(cfg), budget=12, seed=0)
    over = [t for t in rec.trials if t["over_limit"]]
    done = study.get_trials(states=(TrialState.COMPLETE,))
    assert over and len(done) == len(rec.trials)  # every point over the limit taught the model, none failed
    assert all((fz.constraints["requirement"] > 0) == t["over_limit"] for fz, t in zip(done, rec.trials))
    # after the 10 random start-up tries the model keeps to the part under the limit (75% of the space is over it)
    late = rec.trials[10:]
    assert sum(t["over_limit"] for t in late) / len(late) < 0.25  # 0.62 when they were told as failed


def test_check_uses_the_search_memory_metric_and_held_out_ppl_is_never_an_objective(cfg, monkeypatch):
    m = {"model_size_gb": 2.2, "predicted_weight_memory_gb": 0.9}  # a simulated method's 16-bit size vs the plan
    assert search.check_requirement(cfg, m).met and not cfg.requirement.check(m).met
    monkeypatch.setattr(search, "_build_and_measure", lambda *a: {"metrics": {**m, "ppl_val": 10.0}, "errors": {}})
    assert evaluate(start_run(cfg), _point(cfg))["met"]  # after the build too
    with pytest.raises(ValueError, match="memory_metric"):
        run_search(start_run(cfg.with_overrides({"search.memory_metric": "model_size_gb"})))
    with pytest.raises(ValueError, match="ppl_heldout"):
        run_search(start_run(cfg.with_overrides({"search.objectives": ["ppl_heldout", "ppl_val"]})))


def test_mobo_does_not_keep_proposing_a_point_that_fails_after_building(cfg, monkeypatch):
    from optuna.trial import TrialState

    # every H2O model is built and then fails (not cached), the rest are measured
    monkeypatch.setattr(search, "evaluate", _fake_evaluate(over=lambda p: False,
                                                           crash=lambda p: p["stage3_method"] == "h2o" and "building"))
    rec = Recorder(start_run(cfg), "mobo", 6)
    study = run_mobo(rec, technique_space(cfg), budget=20, seed=0)
    assert not study.get_trials(states=(TrialState.FAIL,))  # the failures taught the model
    late = rec.trials[10:]
    assert sum(t["point"]["stage3_method"] == "h2o" for t in late) / len(late) < 0.25
