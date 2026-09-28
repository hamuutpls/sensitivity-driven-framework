"""Excel output for StageReporter: Summary / Config / Per-layer / Raw / Charts."""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from sdf.reporting.metrics import METRICS, label as metric_label

if TYPE_CHECKING:
    from pathlib import Path

    from sdf.reporting.reporter import StageReporter

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BOLD = Font(bold=True)

_DELTA_COLS = (("vs_fp16_abs", "Δ vs FP16"), ("vs_fp16_pct", "Δ% vs FP16"),
               ("vs_original_abs", "Δ vs original"), ("vs_original_pct", "Δ% vs original"))


def metric_order(rep: "StageReporter") -> list[str]:
    present = {k for r in rep.rows for k in r.metrics}
    return [m for m in METRICS if m in present] + sorted(present - set(METRICS))


def write_workbook(rep: "StageReporter", path: "Path") -> None:
    wb = Workbook()
    _summary(wb.active, rep)
    _key_values(wb.create_sheet("Config"), rep)
    _records(wb.create_sheet("Per-layer"), rep.per_layer)
    _records(wb.create_sheet("Raw"), rep.raw)
    _charts(wb.create_sheet("Charts"), rep)
    tmp = path.with_name(f".{path.name}.tmp")
    wb.save(tmp)
    tmp.replace(path)


def _cell(v: Any) -> Any:
    if isinstance(v, (dict, list, tuple)):
        return json.dumps(v, default=str)
    return v


def _header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
    ws.freeze_panes = "C2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(h)) + 2))


def _summary(ws, rep: "StageReporter") -> None:
    ws.title = "Summary"
    metrics = metric_order(rep)
    headers = ["Method", "Variant", "Status", "Requirement met", "Shortfall"]
    col_of: dict[tuple[str, str], int] = {}
    for m in metrics:
        unit = METRICS[m].unit if m in METRICS else ""
        headers.append(m + (f" [{unit}]" if unit else ""))
        for key, label in _DELTA_COLS:
            headers.append(f"{m} {label}")
            col_of[(m, key)] = len(headers)
    headers.append("Error")
    _header(ws, headers)

    for r in rep.rows:
        met = r.requirement.get("met")
        line = [r.method, r.variant, r.status, "n/a" if met is None else ("yes" if met else "no"),
                _cell(r.requirement.get("shortfall") or "")]
        for m in metrics:
            line.append(_cell(r.metrics.get(m)))
            line.extend(r.deltas.get(m, {}).get(key) for key, _ in _DELTA_COLS)
        line.append(r.error or "")
        ws.append(line)

    # Highlight where the framework wins (green) or loses (red) against the original method.
    last = len(rep.rows) + 1
    if last < 2:
        return
    for m in metrics:
        better = METRICS[m].better if m in METRICS else None
        if better is None:
            continue
        win_op, lose_op = ("lessThan", "greaterThan") if better == "lower" else ("greaterThan", "lessThan")
        for key in ("vs_original_abs", "vs_original_pct"):
            col = get_column_letter(col_of[(m, key)])
            rng = f"{col}2:{col}{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator=win_op, formula=["0"], fill=GREEN))
            ws.conditional_formatting.add(rng, CellIsRule(operator=lose_op, formula=["0"], fill=RED))


def _flatten(prefix: str, d: Any, out: list[tuple[str, Any]]) -> None:
    if isinstance(d, dict):
        for k, v in d.items():
            _flatten(f"{prefix}.{k}" if prefix else str(k), v, out)
    else:
        out.append((prefix, _cell(d)))


def _key_values(ws, rep: "StageReporter") -> None:
    _header(ws, ["Setting", "Value"])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 50
    rows: list[tuple[str, Any]] = []
    for section, data in (("conditions", rep.conditions), ("config", rep.config),
                          ("requirement", rep.to_dict()["requirement"]), ("environment", rep.environment)):
        _flatten(section, data, rows)
    for row in rows:
        ws.append(list(row))


def _records(ws, records: list[dict[str, Any]]) -> None:
    if not records:
        ws.append(["(no data for this stage)"])
        return
    headers: list[str] = []
    for r in records:
        headers.extend(k for k in r if k not in headers)
    _header(ws, headers)
    for r in records:
        ws.append([_cell(r.get(h)) for h in headers])


def _charts(ws, rep: "StageReporter") -> None:
    """One bar chart per main metric: FP16 vs original method vs framework."""
    ok = [r for r in rep.rows if r.status == "ok"]
    row0 = 1
    for m in rep.main_metrics:
        points = [(r.key, r.metrics[m]) for r in ok if isinstance(r.metrics.get(m), (int, float))]
        if not points:
            continue
        ws.cell(row=row0, column=1, value="Row").font = BOLD
        ws.cell(row=row0, column=2, value=metric_label(m)).font = BOLD
        for i, (key, value) in enumerate(points, start=1):
            ws.cell(row=row0 + i, column=1, value=key)
            ws.cell(row=row0 + i, column=2, value=value)
        chart = BarChart()
        chart.title = metric_label(m)
        chart.legend = None
        chart.height, chart.width = 7, 14
        chart.add_data(Reference(ws, min_col=2, min_row=row0, max_row=row0 + len(points)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=row0 + 1, max_row=row0 + len(points)))
        ws.add_chart(chart, f"D{row0}")
        row0 += max(len(points) + 2, 16)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
