"""The settings search's outputs in <run_dir>/search/: search_report.md, search_results.xlsx (Trials, Pareto,
Searcher comparison, Best configs), pareto_<objective>_vs_<metric>.png and hypervolume.png."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from sdf.eval import vllm_backend  # noqa: F401 - registers the vLLM metric names
from sdf.reporting.metrics import is_number, label
from sdf.run import RunContext
from sdf.search import FAILED, hypervolume, pareto, tried_path
from sdf.search_space import SearchSpace
from sdf.utils.cache import atomic_write_text
from sdf.utils.logging import get_logger

log = get_logger(__name__)

_SHOWN = ("ppl_val", "ppl_heldout", "predicted_weight_memory_gb", "vllm_decode_ms_per_token", "vllm_ppl_val",
          "vllm_size_gb", "decode_ms_per_token_mean", "avg_bits_per_weight", "avg_activation_bits",
          "decoder_weights_removed", "predicted_kv_memory_gb")


def compare_searchers(trials: list[dict[str, Any]], objectives: list[str]) -> list[dict[str, Any]]:
    """Per searcher: tries, models built (and how many on all the windows), cost, share meeting the requirement, its best first objective and the cost at
    which it first found it, hypervolume (one reference point for all: 10% past the worst met value), and its share
    of the combined trade-off line."""
    met = [t["values"] for t in trials if t["full"] and t["met"]]
    ref = [max(v[i] for v in met) * 1.1 for i in range(len(objectives))] if met else None
    front = {json.dumps(t["point"], sort_keys=True) for t in pareto(trials)}
    out = []
    for name in dict.fromkeys(t["searcher"] for t in trials):
        ts = [t for t in trials if t["searcher"] == name]
        full_met = [t for t in ts if t["full"] and t["met"]]
        best = min(full_met, key=lambda t: t["values"][0], default=None)
        mine = {json.dumps(t["point"], sort_keys=True) for t in full_met}
        out.append({
            "Search method": SEARCHER_NAMES.get(name, name), "Tries": len(ts),
            "Models built": sum(1 for t in ts if t["cost"] > 0 and not t.get("earlier_run")),
            "Of those, on all evaluation windows": sum(1 for t in ts if t["cost"] > 0 and t["full"]
                                                       and not t.get("earlier_run")),
            "Cost, in full tries": round(sum(t["cost"] for t in ts), 3),
            "Time, minutes": round(sum(t.get("time_s", 0.0) for t in ts) / 60, 1),
            "Share of tries under the limit (higher is better)": len([t for t in ts if t["met"]]) / len(ts),
            f"Best {label(objectives[0]).lower()} (lower is better)": best["values"][0] if best else None,
            "Cost until it found that best (lower is better)":
                round(sum(t["cost"] for t in ts[: ts.index(best) + 1]), 3) if best else None,
            "Hypervolume (higher is better)": hypervolume(ts, ref) if ref else 0.0,
            "Share of the best trade-off line it found (higher is better)":
                len(front & mine) / len(front) if front else 0.0,
        })
    return out


SEARCHER_NAMES = {"mobo": "MOBO (multi-objective Bayesian optimisation)",
                  "mfbo": "MFBO (multi-fidelity: rough measurements first, full ones for the best)",
                  "nsga3": "NSGA-III (an evolutionary method)"}
SETTING_NAMES = {"sensitive_threshold": "Protection threshold", "prune_ratio_aggressive": "Pruning share",
                 "calib_dataset": "Calibration text", "calib_samples": "Calibration passages",
                 "gptq_groupsize": "Group size", "smoothquant_alpha": "SmoothQuant strength",
                 "quarot_k_bits": "QuaRot key bits", "stage1_method": "Stage 1 quantization",
                 "stage2_method": "Stage 2 pruning",
                 "stage3_method": "Stage 3 KV-cache (memory of earlier words) technique"}


def setting_value(name: str, value: Any) -> str:
    """A setting in words: techniques by their label, numbers rounded."""
    from sdf.stages.methods import METHODS
    from sdf.stages.runner import stage1_row

    if name == "stage1_method":
        return stage1_row(value)[0].label
    if name in ("stage2_method", "stage3_method"):
        return METHODS[value].label
    return _fmt(value)


def hypervolume_curve(trials: list[dict[str, Any]], searcher: str, ref: list[float]) -> list[tuple[float, float]]:
    """(cost spent, hypervolume so far) after each of `searcher`'s evaluations."""
    ts = [t for t in trials if t["searcher"] == searcher]
    return [(sum(t["cost"] for t in ts[: i + 1]), hypervolume(ts[: i + 1], ref)) for i in range(len(ts))]


def _backend_rows(backends: dict[str, Any]) -> list[list[str]]:
    """One row per backend measure, each per-token latency followed by its tokens per second (report rules)."""
    rows = []
    for m, v in backends.items():
        rows.append([label(m), _fmt(v)])
        if m.endswith("_decode_ms_per_token") and is_number(v) and v > 0 and \
                m.replace("_ms_per_token", "_tokens_per_s") not in backends:
            rows.append([f"{label(m)}, as tokens per second", _fmt(1000.0 / v)])
    return rows


def write_search_report(ctx: RunContext, trials: list[dict[str, Any]], space: SearchSpace,
                        intro: str | None = None, extras: list[dict[str, Any]] = (),
                        accuracy: dict[str, Any] | None = None) -> dict[str, Path]:
    """`intro`: the report's opening paragraph when the search is not the chained one (a per-stage search).
    `extras`: tries made apart from the searchers (sdf.search.extra_tries), in a section of their own. `accuracy`:
    the Stage 4 test of the final model best on the first objective, backends and task accuracy (sdf.search.final_model)."""
    cfg = ctx.cfg
    s = cfg.search
    d = ctx.run_dir / "search"
    d.mkdir(parents=True, exist_ok=True)
    comparison = compare_searchers(trials, s.objectives)
    front = sorted(pareto(trials), key=lambda t: t["values"][0])
    best = {o: min((t for t in trials if t["full"] and t["met"]), key=lambda t, i=i: t["values"][i], default=None)
            for i, o in enumerate(s.objectives)}
    names = [p.name for p in space.params]
    outputs = {"search_trials": d / "trials.json"}
    atomic_write_text(outputs["search_trials"], json.dumps(trials, indent=2, default=str))

    # --- workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Trials"
    settings = [SETTING_NAMES.get(n, n) for n in names]
    results = [label(m) for m in _SHOWN]
    values = lambda t: [setting_value(n, t["point"][n]) for n in names] + [t["metrics"].get(m) for m in _SHOWN]  # noqa: E731
    _table(ws, ["Search method", "Try", "Evaluation windows", "Cost, in full tries", "Time, s", "Under the limit",
                "Why it failed"] + settings + results,
           [[t["searcher"], t["n"], t["windows"], t["cost"], t.get("time_s"), t["met"], t["error"]] + values(t)
            for t in trials])
    _table(wb.create_sheet("Pareto"), ["Search method", "Try"] + settings + results,
           [[t["searcher"], t["n"]] + values(t) for t in front])
    cols = list(comparison[0]) if comparison else ["Search method"]
    _table(wb.create_sheet("Searcher comparison"), cols, [[c[k] for k in cols] for c in comparison])
    _table(wb.create_sheet("Best configs"), ["Best for"] + settings + results,
           [[label(o)] + values(t) for o, t in best.items() if t])
    outputs["search_xlsx"] = d / "search_results.xlsx"
    wb.save(outputs["search_xlsx"])

    outputs.update(_plots(d, trials, front, s.objectives, s.memory_metric, comparison))

    # --- report
    shown = list(dict.fromkeys(m for m in (s.objectives[0], s.memory_metric, *s.objectives[1:], "ppl_heldout") if m))
    tech = [n for n in ("stage1_method", "stage2_method", "prune_ratio_aggressive", "stage3_method") if n in names]
    rest = [n for n in names if n not in tech]
    lines = [
        "# Settings search", "",
        "## What this shows", "",
        "The framework tried many settings and built each one as a chained model: the Stage 0 plan (which layers are "
        "sensitive), then Stage 1 (weights and activations stored with fewer bits, following the joint weight and activation "
        "plan), then Stage 2 (pruning: removing "
        "numbers), then Stage 3 (compressing the KV cache, the model's memory of earlier words). A setting is the "
        "technique of each stage plus the numbers that tune them (protection threshold, pruning share, calibration "
        f"text and so on). Only settings that meet the limit count ({_req(cfg)}). Among those the search looks for "
        "the lowest " + " and the lowest ".join(label(o).lower() for o in s.objectives) + ". Speed comes from vLLM, "
        "which runs the model up to Stage 2 (it cannot run the Stage 3 techniques); unstructured pruning keeps the "
        "removed numbers as zeros, so it does not make vLLM faster. "
        + ("Stage 0's refinement was off: every plan is the one-shot joint plan. "
           if cfg.stage0.joint_refine_rounds == 0 else "")
        + f"{len(s.searchers)} search method{'s' * (len(s.searchers) != 1)} each had the same budget of "
        f"{s.budget} full tries, where a full try is one model measured on all the evaluation windows. MFBO's rough "
        "tries use fewer windows and count by that share, so it builds more models for the same budget (see Models "
        "built below). A setting whose plan is already over the limit is not built and does not count, nor does one "
        "that failed before building or a repeat of the method's own try that came from the cache. The methods: "
        + "; ".join(SEARCHER_NAMES[n] for n in s.searchers if n in SEARCHER_NAMES) + "." if intro is None else intro, "",
        "## What it means", "",
        f"{len(front)} setting(s) are on the best trade-off line: none of them can get better on one goal without "
        "getting worse on another, so each is a fair choice depending on what matters more. They are listed below, "
        f"best {label(s.objectives[0]).lower()} first (lower is better on every goal). Held-out perplexity "
        "was measured but never used to choose, so it shows how well the choice holds on unseen text.", "",
        "## Best trade-off settings", "",
        _md([["#"] + [{"stage3_method": "Stage 3 technique"}.get(n, SETTING_NAMES[n]) for n in tech]
             + [label(m) + " (lower is better)" for m in shown]]
            + [[str(i + 1)] + [setting_value(n, t["point"][n]) for n in tech] + [_fmt(t["metrics"].get(m)) for m in shown]
               for i, t in enumerate(front)]),
        "The other settings of each option:", "",
        _md([["#"] + [SETTING_NAMES.get(n, n) for n in rest]]
            + [[str(i + 1)] + [setting_value(n, t["point"][n]) for n in rest] for i, t in enumerate(front)]),
        "## Search methods compared", "",
        _md([list(comparison[0]) if comparison else ["Search method"]]
            + [[_fmt(v) for v in c.values()] for c in comparison]),
        "A full try measures on every evaluation window (a slice of the test text); MFBO's rough tries use fewer "
        "windows and count by that share, though each still builds the whole model, so its time column is the fairer "
        "cost. Hypervolume measures how much of the space of good results (low "
        + " and low ".join(label(o).lower() for o in s.objectives) + ") the method's settings under the limit cover; "
        "larger is better.", "",
        "## Conditions", "",
        f"- Model: {cfg.model.name}; seed {cfg.run.seed}; evaluation windows: {cfg.eval.max_windows or 'all'} of "
        f"{cfg.eval.seq_len} tokens.",
        "- Techniques allowed: Stage 1 " + ", ".join(setting_value("stage1_method", m) for m in s.stage1_methods)
        + "; Stage 2 " + ", ".join(setting_value("stage2_method", m) for m in s.stage2_methods)
        + "; Stage 3 " + ", ".join(setting_value("stage3_method", m) for m in s.stage3_methods) + ".",
        f"- Tries over the limit, so not built: {sum(1 for t in trials if t.get('over_limit'))} of {len(trials)}.",
        f"- Tries whose model an earlier run already built, so not built again (its earlier result is used): "
        f"{sum(1 for t in trials if t.get('earlier_run') not in (None, ctx.run_id))} of {len(trials)}. Tries that "
        f"were the same model as an earlier try of this run (they differ only in a setting their techniques do not "
        f"read), so not built either: {sum(1 for t in trials if t.get('earlier_run') == ctx.run_id)}. Every "
        f"combination tried so far is listed in {tried_path(ctx)}.",
        f"- Tries that failed (could not be built or measured): "
        f"{sum(1 for t in trials if not t['ok'] and not t.get('over_limit'))} of "
        f"{len(trials)}; the reasons are in the Trials sheet of search_results.xlsx.", "",
    ]
    if accuracy:
        acc = {m: v for m, v in (accuracy["metrics"] or {}).items() if m.startswith(("acc_", "downstream_acc"))}
        backends = {m: v for m, v in (accuracy["metrics"] or {}).items() if m.startswith(("vllm_", "llamacpp_", "trtllm_"))}
        pick = f"the final model under the limit with the best {label(s.objectives[0]).lower()}"
        which = (f"the extra {setting_value('stage1_method', accuracy['point']['stage1_method'])} try"
                 if accuracy["searcher"] == "extra" else f"{accuracy['searcher'].upper()} try #{accuracy['n']}") + (
            f" of the earlier run {accuracy['earlier_run']}, a combination this run did not try again"
            if accuracy.get("earlier_run") else "")
        lines += [
            "## The final model on the backends (Stage 4)", "",
            f"The final model is {which}: " + ", ".join(
                f"{label(m).lower()} {_fmt(v)}" for m, v in (accuracy.get("try_metrics") or {}).items()
                if m in ("predicted_weight_memory_gb", "ppl_val", "ppl_heldout")) + "; settings: " + ", ".join(
                f"{SETTING_NAMES.get(k, k).lower()} {setting_value(k, v)}" for k, v in sorted((accuracy.get("point") or {}).items())) + ".", "",
            f"**What this shows:** {pick} ({which}) measured on each serving backend; the tries "
            "themselves ran on Hugging Face only. **What it means:** how fast and how accurate the chosen model is "
            "where it would be deployed. A backend re-rounds the model into its own file format, so its perplexity "
            "is that file's, not the verdict.", "",
            _md([["Measure", "Value"]] + _backend_rows(backends))
            if backends else (f"The final model could not be built again: {accuracy['error']}." if accuracy.get("error")
                              else "No backend result: " + ("; ".join(f"{k} failed: {v}" for k, v in
                                                                      accuracy["errors"].items())
                                                            if accuracy.get("errors") else "no backend was set") + "."),
            "",
            "## Accuracy of the final model", "",
            "**What this shows:** how many test questions the best final model still answers correctly, next to "
            "the uncompressed TinyLlama. **What it means:** a small change says the compression kept the model's "
            "ability; a large drop says the lower perplexity did not carry over to answering questions.", "",
            f"{pick[0].upper() + pick[1:]} ({which}) was built again and tested on multiple-choice tasks: the share of questions it "
            "answers correctly (higher is better), next to the uncompressed model's on the same tasks, measured in the "
            "same run.", "",
            _md([["Task", "Final model", "Uncompressed model", "Change", "Change, %"]]
                + [[label(m), _fmt(v), _fmt(base := accuracy["fp16"].get(m)),
                    _fmt(v - base) if base is not None else "",
                    _fmt(100 * (v - base) / base) if base else ""] for m, v in acc.items()])
            if acc else f"No accuracy result: {accuracy.get('error') or 'no task set, or the test failed'}.", ""]
    if extras:
        lines += [
            "## Extra Stage 1 techniques", "",
            "**What this shows:** SpQR, QuIP and QuIP# (the techniques asked for) when no search method built them: "
            "one chained try each, the best setting under the limit with only its Stage 1 technique changed. "
            "**What it means:** how each would have done in the best model; they are not part of the comparison or "
            "the trade-off line.", "",
            _md([["Stage 1 technique", "Our method WxAy (average bits)", "Under the limit"]
                 + [label(m) + " (lower is better)" for m in shown] + ["Why it failed"]]
                + [[setting_value("stage1_method", t["point"]["stage1_method"]), _wxay(t["metrics"]),
                    "yes" if t["met"] else "no"]
                   + [_fmt(t["metrics"].get(m)) for m in shown] + [str(t["error"] or "")] for t in extras])]
    outputs["search_report"] = d / "search_report.md"
    atomic_write_text(outputs["search_report"], "\n".join(lines))
    return outputs


def _req(cfg) -> str:
    r = cfg.requirement
    parts = [f"{label(m)} at most {v}" for m, v in ((cfg.search.memory_metric, r.target_memory_gb),
                                                     ("ppl_val", r.target_ppl),
                                                     ("decode_ms_per_token_mean", r.target_latency_ms),
                                                     ("predicted_kv_memory_gb", r.kv_budget_gb)) if v is not None]
    return "; ".join(parts) or "no limits set"


def _wxay(metrics: dict[str, Any]) -> str:
    """The row's average weight and activation bits as "W4.5A8", or "" when not measured."""
    w, a = metrics.get("decoder_weight_bits"), metrics.get("avg_activation_bits")
    return f"W{w:.3g}A{a:.3g}" if isinstance(w, (int, float)) and isinstance(a, (int, float)) else ""


def _fmt(v: Any) -> str:
    if isinstance(v, float):
        return "" if v >= FAILED else f"{v:.4g}"
    return "" if v is None else str(v)


def _md(rows: list[list[str]]) -> str:
    head, *body = rows
    return "\n".join(["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
                     + ["| " + " | ".join(r) + " |" for r in body]) + "\n"


def _table(ws, header: list[str], rows: list[list[Any]]) -> None:
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([v if isinstance(v, (int, float, str, bool)) or v is None else str(v) for v in r])


def _plots(d: Path, trials, front, objectives, memory, comparison) -> dict[str, Path]:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        log.warning("matplotlib is not installed: no search plots (pip install .[figures])")
        return {}
    out = {}
    ok = [t for t in trials if t["full"] and t["ok"]]
    colors = dict(zip(dict.fromkeys(t["searcher"] for t in trials), ("#0072B2", "#E69F00", "#009E73", "#CC79A7")))
    for x in [o for o in dict.fromkeys([memory, *objectives]) if o != "ppl_val"]:  # each against perplexity
        fig, ax = plt.subplots(figsize=(6, 4))
        for name, c in colors.items():
            pts = [t for t in ok if t["searcher"] == name and t["metrics"].get(x) is not None]
            ax.scatter([t["metrics"][x] for t in pts], [t["metrics"]["ppl_val"] for t in pts], c=c, label=name,
                       alpha=0.6, marker="o")
        pts = [t for t in front if t["metrics"].get(x) is not None]
        ax.scatter([t["metrics"][x] for t in pts], [t["metrics"]["ppl_val"] for t in pts], facecolors="none",
                   edgecolors="black", s=120, label="best trade-off")
        ax.set_xlabel(label(x) + " (lower is better)")
        ax.set_ylabel("validation perplexity (lower is better)")
        ax.legend()
        fig.tight_layout()
        out[f"search_pareto_{x}"] = d / f"pareto_ppl_vs_{x}.png"
        fig.savefig(out[f"search_pareto_{x}"], dpi=150)
        plt.close(fig)
    met = [t["values"] for t in ok if t["met"]]
    if met:
        ref = [max(v[i] for v in met) * 1.1 for i in range(len(objectives))]
        fig, ax = plt.subplots(figsize=(6, 4))
        for name, c in colors.items():
            curve = hypervolume_curve(trials, name, ref)
            ax.step([a for a, _ in curve], [b for _, b in curve], where="post", c=c, label=name)
        ax.set_xlabel("cost (full evaluations)")
        ax.set_ylabel("hypervolume (higher is better)")
        ax.legend()
        fig.tight_layout()
        out["search_hypervolume"] = d / "hypervolume.png"
        fig.savefig(out["search_hypervolume"], dpi=150)
        plt.close(fig)
    return out
