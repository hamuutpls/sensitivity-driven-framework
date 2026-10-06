"""End of a full run: all_stages_comparison.xlsx and master_report.md from every stage_<N>/results.json.

Built from the saved JSON only, so it can be (re)written at any time, also after an interrupted run:
    python -c "from sdf.reporting.master import write_master; write_master('thesis_compression/results/<run_id>')"
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

from sdf.reporting.excel import GREEN, RED, _header, save_workbook
from sdf.reporting.markdown import _column_notes, _table, original_model_lines, requirement_cell
from sdf.reporting.metrics import METRICS, is_number, label
from sdf.utils.cache import atomic_write_text

# Columns of the cross-stage table, when a stage has them. Every one has a better direction in METRICS.
COLUMNS = ["ppl_val", "ppl_heldout", "downstream_acc_mean", "predicted_weight_memory_gb", "model_size_gb",
           "avg_activation_bits", "predicted_kv_memory_gb", "peak_memory_gb", "decode_ms_per_token_mean",
           "build_time_s"]


def load_stages(run_dir: str | Path) -> list[dict[str, Any]]:
    """results.json of every stage_<N>/ folder in the run, by stage number (side studies are left out)."""
    out = []
    for p in Path(run_dir).glob("stage_*/results.json"):
        if p.parent.name.removeprefix("stage_").isdigit():
            out.append(json.loads(p.read_text(encoding="utf-8")))
    return sorted(out, key=lambda r: r["stage"])


def _rows(stages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for st in stages:
        for r in st["rows"]:
            rows.append({"stage": st["stage"], "row": f"{r['method']}/{r['variant']}", "status": r["status"],
                         "metrics": r["metrics"], "deltas": r.get("deltas", {}), "requirement": r.get("requirement", {}),
                         "error": r.get("error")})
    return rows


def write_master(run_dir: str | Path) -> dict[str, Path]:
    run_dir = Path(run_dir)
    stages = load_stages(run_dir)
    if not stages:
        raise FileNotFoundError(f"no stage_<N>/results.json under {run_dir}")
    rows = _rows(stages)
    cols = [c for c in COLUMNS if any(c in r["metrics"] for r in rows)]
    xlsx = run_dir / "all_stages_comparison.xlsx"
    _workbook(stages, rows, cols, xlsx)
    md = run_dir / "master_report.md"
    atomic_write_text(md, _markdown(run_dir, stages, rows, cols))
    return {"master_report": md, "all_stages_xlsx": xlsx}


def _workbook(stages, rows, cols, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "All stages"
    headers = ["Stage", "Row", "Status"] + [label(c) for c in cols] + ["Δ% perplexity (val) vs original",
                                                                      "Requirement met"]
    _header(ws, headers)
    for r in rows:
        pct = r["deltas"].get("ppl_val", {}).get("vs_original_pct")
        ws.append([r["stage"], r["row"], r["status"]] + [r["metrics"].get(c) for c in cols]
                  + [pct, requirement_cell(r["requirement"])])
    # Win/loss vs the standard method: lower perplexity is better, so a negative change is a win.
    col = get_column_letter(len(headers) - 1)
    rng = f"{col}2:{col}{len(rows) + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=RED))

    notes = wb.create_sheet("Columns")
    notes.append(["Column", "Better", "Meaning"])
    for c in cols:
        notes.append([label(c), METRICS[c].better or "", METRICS[c].meaning])
    notes.append(["Δ% perplexity (val) vs original", "lower",
                  "Change in validation perplexity of a framework row against its method's standard row; "
                  "green = framework more accurate, red = less."])

    fs = wb.create_sheet("Findings")
    fs.append(["Stage", "Finding"])
    for st in stages:
        for f in st.get("findings", []):
            fs.append([st["stage"], f])
    save_workbook(wb, path)


def _cell(v: Any, metric: str) -> str:
    if not is_number(v):
        return "–"
    return f"{v:.4g} {METRICS[metric].unit}".strip()


def _markdown(run_dir: Path, stages, rows, cols) -> str:
    first = stages[0]
    lines = [f"# Master report: run {run_dir.name}", "",
             "One page over every stage of this run. Each stage compares the uncompressed model, the standard "
             "compression method and the sensitivity-guided framework under identical conditions; the details are "
             "in each stage's own report, linked below.", ""]
    lines += original_model_lines(first.get("original_model", {}))
    lines += ["## Stages", ""]
    for st in stages:
        lines.append(f"### Stage {st['stage']}: {st['title']} ([report](stage_{st['stage']}/report.md))")
        lines.append("")
        lines += [f"- {f}" for f in st.get("findings", [])] or ["- No findings recorded."]
        failed = [r for r in st["rows"] if r["status"] != "ok"]
        if failed:
            lines.append(f"- {len(failed)} row(s) failed: "
                         + "; ".join(f"`{r['method']}/{r['variant']}` ({r.get('error')})" for r in failed))
        lines.append("")
    lines += ["## Every row", ""]
    headers = ["Stage", "Row"] + [f"{label(c)} ({METRICS[c].better} is better)" for c in cols] + ["Requirement met"]
    table = [[r["stage"], f"`{r['row']}`"] + [_cell(r["metrics"].get(c), c) for c in cols]
             + [requirement_cell(r["requirement"])] for r in rows if r["status"] == "ok"]
    lines += [_table(headers, table), ""]
    lines += _column_notes([("Row", "Method and version: `fp16` uncompressed, `original` the standard method, "
                                    "`framework` guided by Stage 0.")]
                           + [(label(c), METRICS[c].meaning) for c in cols])
    lines.append("All rows with deltas and the findings are in `all_stages_comparison.xlsx`.")
    return "\n".join(lines) + "\n"
